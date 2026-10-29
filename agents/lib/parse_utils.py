#!/usr/bin/env python3
"""Shared PDF/Excel/CSV parsing utilities for tax document extraction.

Used by /tax-parse subagents to extract text and tables from source documents.
Wraps pdfplumber with error handling and provides consistent interfaces.
"""

import bisect
import csv
import functools
import json
import math
import re
import sys
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
from agents.lib import preflight  # noqa: E402
from agents.lib.secure_fs import (  # noqa: E402
    identifier_spans, is_reparse_point, mask_row_runs, mask_runs, mask_ssn, mask_table, names_an_identifier,
    prepare_runs, redact)
from agents.lib.workspace import WorkspaceError, resolve_workspace  # noqa: E402

# Activate defusedxml to protect openpyxl against XXE attacks in crafted .xlsx files.
# Must be called before any XML parsing (including openpyxl imports).
#
# The OUTCOME is recorded, not just announced. A warning on stderr is not a
# guard: this module went on to open workbooks with the stdlib XML parser
# restored, on the one input class AGENTS.md §7 names defusedxml to defend —
# an untrusted `.xlsx` a taxpayer downloaded from a payer. Worse, it is the
# failure the preflight gate reasons FROM: that gate warns rather than blocks on
# a missing package because "it surfaces loudly at first use, as an ImportError
# inside /tax-parse", which was true of the other three and false of this one
# alone. `extract_excel` refuses instead, so the claim holds for all four.
#
# The refusal is at the Excel entry point rather than at import, because PDF and
# CSV extraction are unaffected and this module is what /tax-parse, /tax-verify
# and /tax-parse-crypto all import first.
try:
    import defusedxml
    defusedxml.defuse_stdlib()
    _XXE_GUARD_ACTIVE = True
except ImportError:
    _XXE_GUARD_ACTIVE = False
    print(
        "WARNING: defusedxml not installed — .xlsx parsing is vulnerable to XXE attacks. "
        "Install with: pip install defusedxml",
        file=sys.stderr,
    )


MAX_FILE_SIZE = 100 * 1024 * 1024  # 100 MB — prevents DoS via oversized files

# V11 — the budgets the on-disk cap above does NOT bound. Source documents are
# untrusted input (AGENTS.md §1), and 100 MB of `.xlsx` is 100 MB of ZIP: the
# XML inside it can be orders of magnitude larger, and neither the extracted
# text nor the time spent producing it is bounded by the file's size at all.
# Every one of these refuses rather than truncating — a partial extraction that
# looks like a complete one is what the tie-out gate cannot see.
#
# The numbers are deliberately far above any real return and are stated here
# rather than passed in: a budget a caller can raise is one the caller raises.
# A consolidated brokerage 1099 runs to a few hundred pages and a few hundred
# kilobytes of text; a workbook's XML is a handful of times its ZIP.
MAX_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024      # 1 GB of XML behind a workbook
MAX_COMPRESSION_RATIO = 200                       # 100 KB expanding to 20 MB
MAX_EXTRACTED_CHARS = 20 * 1024 * 1024            # 20 MB of returned text
MAX_PAGES = 2000                                  # pages in one PDF
MAX_OVERLAY_INDEX_ENTRIES = 100_000              # per-page spatial-index memory budget
PARSE_TIME_BUDGET_SECONDS = 300                   # wall clock inside one call


class ParseBudgetError(ValueError):
    """A document exceeded a parser budget. A `ValueError`, because every
    caller already handles the size cap's refusal that way and a new exception
    class here would slip past those handlers as a traceback."""


class _Budget:
    """The output and time budgets for ONE extraction call.

    Checked incrementally — between pages, sheets and rows — because the point
    is to stop a document that is still expanding, and a check that only runs
    at the end has already paid the cost it exists to avoid.
    """

    def __init__(self, what: str):
        self.what = what
        self.chars = 0
        self.started = time.monotonic()

    def charge_cell(self, value) -> None:
        """Charge one cell: its text as the result carries it, plus the
        punctuation the JSON around it costs. `None` and a number are cells
        too — they are in the returned rows and in the printed output."""
        text = value if isinstance(value, str) else str(value)
        self.charge(text + " " * _CELL_OVERHEAD)

    def tick(self) -> None:
        """The TIME half, checked at every row and page boundary whether or not
        anything was produced.

        Separate from `charge` because the two are bypassed differently: a
        numeric-only worksheet, a blank CSV row and a PDF page with no tables
        all produce no CHARACTERS, and a deadline checked only when a string
        arrives is a deadline those documents never reach.
        """
        if time.monotonic() - self.started > PARSE_TIME_BUDGET_SECONDS:
            raise ParseBudgetError(
                f"{self.what} is still running after "
                f"{PARSE_TIME_BUDGET_SECONDS}s. Refusing rather than returning "
                "a partial extraction: split the file and parse the part you "
                "need.")

    def charge(self, produced) -> None:
        """Charge a produced STRING — a page's text."""
        self.chars += len(produced) if isinstance(produced, str) else 0
        if self.chars > MAX_EXTRACTED_CHARS:
            raise ParseBudgetError(
                f"{self.what} produced more than "
                f"{MAX_EXTRACTED_CHARS // (1024 * 1024)} MB of text, which no "
                "tax document does. Refusing rather than returning a partial "
                "extraction: split the file and parse the part you need.")
        self.tick()


# What a cell costs beyond its own text: the quotes, the comma and the space
# the CLI's JSON puts around it. Small, and the point is that it is not ZERO —
# a numeric or empty cell costs no CHARACTERS and still occupies memory and
# prints, so charging strings alone let a 100x4 numeric sheet through a budget
# it was twenty-six times over.
_CELL_OVERHEAD = 2


def _charge_deep(budget: "_Budget", node) -> None:
    """Charge every CELL in a nested result, and every CONTAINER around them.

    A container costs its own brackets, which is not pedantry: an empty row is
    `[]`, it charged nothing while still being retained and returned, and a
    file of nothing but newlines allocated a list per line under a budget that
    never noticed.
    """
    if isinstance(node, dict):
        budget.charge_cell("")
        for value in node.values():
            _charge_deep(budget, value)
    elif isinstance(node, (list, tuple)):
        budget.charge_cell("")
        for value in node:
            _charge_deep(budget, value)
    else:
        budget.charge_cell(node)


def _refuse_a_zip_bomb(path: str) -> None:
    """Refuse a workbook whose XML is far larger than its ZIP.

    `.xlsx` is a ZIP, and the 100 MB cap above bounds the ZIP alone: a few
    hundred kilobytes of highly compressible XML expands to gigabytes, and
    openpyxl decompresses it into memory before this module sees a single cell.
    Both the absolute size and the RATIO are checked — a bomb is small on disk,
    which is the whole trick.

    RESIDUAL, stated rather than implied: the sizes come from the ZIP's own
    central directory, which a crafted file can understate. The output budget
    downstream is what catches that; this check is what stops the ordinary bomb
    before anything is decompressed at all.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, OSError) as e:
        raise ParseBudgetError(
            f"{redact(path)} is not a readable workbook: {redact(e)}") from e
    uncompressed = sum(entry.file_size for entry in entries)
    compressed = sum(entry.compress_size for entry in entries) or 1
    if uncompressed > MAX_UNCOMPRESSED_BYTES:
        raise ParseBudgetError(
            f"{redact(path)} declares {uncompressed // (1024 * 1024)} MB of "
            f"content behind a {compressed // 1024} KB archive, over the "
            f"{MAX_UNCOMPRESSED_BYTES // (1024 * 1024)} MB budget. Refusing to "
            "decompress it.")
    if uncompressed // compressed > MAX_COMPRESSION_RATIO:
        raise ParseBudgetError(
            f"{redact(path)} expands {uncompressed // compressed}x, over the "
            f"{MAX_COMPRESSION_RATIO}x budget — a real workbook does not. "
            "Refusing to decompress it.")


def _refuse_too_many_pages(requested: int, path: str) -> None:
    """Bound what THIS CALL will extract, not what the file holds.

    /tax-parse splits a long document into per-page sub-tasks (step 2.5's
    page-count probe is what decides), and a cap on the document would refuse
    every one of those slices for a file the workflow is already handling
    correctly. What has to be bounded is the work one call does.
    """
    if requested > MAX_PAGES:
        raise ParseBudgetError(
            f"{redact(path)}: this call would extract {requested} pages, over "
            f"the {MAX_PAGES}-page budget. Refusing rather than returning a "
            "partial extraction: ask for the pages you need (the `pages` "
            "argument) or split the file.")


# Every string this module returns — page text, a table or workbook cell, a
# CSV cell, a sheet name — passes through a masker before it leaves the
# function that read it; see `secure_fs.mask_ssn` for the shape and why it is
# not `redact`. The CLI below is a plain printer of those returns.
#
# TABLES go through `secure_fs.mask_table` rather than through `mask_ssn` cell
# by cell, because the label that makes a bare nine-digit run an identifier
# sits in a NEIGHBOURING cell — a column header, or the cell to its left — and
# a per-cell masker cannot see it. A `_mask_deep` that walked the rows
# individually is what this replaced, and a CSV with an `SSN` column came
# straight back out of it.


# Every distribution `requirements.txt` pins that a document parse actually runs
# on — pdfplumber and the stack under it, openpyxl and the stack under it. The
# lock exists for exactly this set: pdfplumber 0.11.10 shipped no pdfplumber
# code at all, only a `pdfminer.six` bump (AGENTS.md §7), so the package whose
# republication changes what a W-2 parses to is not one anybody names directly.
#
# Split rather than "the whole lock", because a gate that fires on a package the
# parse never reaches is the cry-wolf shape this repository has been bitten by:
# refusing to read a 1099 because the PDF *generator* was upgraded teaches an
# agent that these refusals are noise. The two sets are pinned as a PARTITION of
# the lock by `test_every_locked_distribution_is_classified`, so a bump that
# pulls in a new transitive package fails there — the alternative is that it
# lands in neither set and silently escapes the check.
_PARSING_TREE = frozenset({
    "pdfplumber", "pdfminer-six", "pypdfium2", "pillow", "charset-normalizer",
    "cryptography", "cffi", "pycparser",
    "openpyxl", "et-xmlfile", "defusedxml",
})
# Pinned, and reached by no parse: `reportlab` writes the summary PDF in
# /tax-generate, downstream of every figure this module produces.
_GENERATION_TREE = frozenset({"reportlab"})


@functools.lru_cache(maxsize=1)
def _off_lock_parsers() -> str:
    """Why the parsing stack is not the locked one, or "" when it is.

    Cached: this is read at every PDF and workbook entry point, and it walks the
    installed distributions and re-reads the lock. A parse splits into per-
    document sub-tasks, so the walk would otherwise repeat once per file, and
    within one process the answer cannot change. Tests that fake an environment
    call `cache_clear()`.
    """
    drift, lock_error = preflight._drifted_dependencies()
    named = [record for record in drift if record[0] in _PARSING_TREE]
    if named:
        # `installed is None` is the ABSENT case, carried as a state rather than
        # as prose — printing the sentinel would put "pdfminer-six None" in
        # front of the reader instead of the finding.
        detail = ", ".join(
            f"{pkg} is NOT INSTALLED (lock pins {pinned})" if installed is None
            else f"{pkg} {installed} (lock pins {pinned})"
            for pkg, installed, pinned in named)
        return ("the installed parsing packages are not the ones the lock "
                f"pins: {detail}")
    if lock_error:
        # Unverifiable, not verified-clean. `requirements.txt` is tracked and
        # sits two levels above this file, so a workspace that cannot read it is
        # one where `agents/lib/` was moved away from the clone it belongs to —
        # the state AGENTS.md §4 says must not compute a return, arrived at from
        # the other side.
        return ("the installed parsing packages could not be compared against "
                f"requirements.txt at all: {lock_error}")
    return ""


def _require_locked_parsers() -> None:
    """Refuse to open a document while the parsing stack is off the lock.

    The preflight gate WARNS about the same records, and the two are not in
    conflict — they are the two halves of one decision. Blocking the gate would
    refuse /tax-calendar and every other skill that opens no document, on a
    signal parsed out of a text file. Refusing HERE costs those skills nothing
    and covers the case the lock is actually for: a parser the suite never ran
    against, reading a taxpayer's document, producing a figure that is wrong in
    a direction nobody can predict and that no downstream check can catch —
    `tax-review` audits the aggregated result, and the tie-out compares the
    ledger to the result, so a misparsed box ties out perfectly.
    """
    off_lock = _off_lock_parsers()
    if off_lock:
        raise RuntimeError(
            f"{off_lock}. A parser version the lock does not name is one the "
            "test suite never ran against, and a republished parser changes "
            "what a document parses to — which is why requirements.txt pins "
            "the whole transitive tree by hash (AGENTS.md §7). Source "
            "documents are untrusted input and the return's figures come off "
            "this parse, so this is a refusal, not a warning. Restore the "
            "pinned tree: "
            "sh agents/lib/run.sh -m pip install -r requirements.txt")


def _assert_within(candidate, workspace: str | None, *,
                   subject: str, verb: str) -> Path:
    """`candidate` resolved, or `ValueError` if it lands outside `workspace`.

    ONE containment predicate. The question — does this resolved path sit under
    this resolved base — was answered at three sites in this file, and the third
    copy had already drifted: it resolved and compared without the fail-closed
    reparse pre-check the other two sit behind. A containment fix that has to
    land three times lands twice.

    The subject and verb are the caller's, because "path" and "symlink" are
    different findings to whoever reads the refusal — a link POINTS outside,
    which says the file was reached through one, and that distinction is what
    the existing assertions match on. Both sides are redacted: a workspace named
    after a taxpayer discloses one.
    """
    resolved = Path(candidate).resolve()
    if workspace is None:
        return resolved
    if not resolved.is_relative_to(Path(workspace).resolve()):
        raise ValueError(
            f"{subject} '{redact(candidate)}' {verb} outside workspace "
            f"'{redact(workspace)}'")
    return resolved


def _validate_file_path(
    path: str,
    expected_suffixes: set[str] | None = None,
    workspace: str | None = None,
) -> Path:
    """Validate that a file path exists, check extension, and enforce containment.

    Checks reparse points on the original (unresolved) path before resolving,
    then verifies the resolved path stays within the workspace.

    **Callers open the RETURNED path, never the string they passed in.** All five
    used to re-open the raw argument, so the containment and suffix checks ran on
    the resolved path while the read ran on the unresolved one — a link swapped
    between the two was read unchecked. The race needs local write access to the
    directory and is not reachable from a document or from the parse worker, but
    closing it costs one assignment.

    RESOLVE DOES NOT SEE A HARD LINK, and that is a deliberate residual: a hard
    link planted in `documents/` points at another file's inode with no path to
    resolve, so it reads as an ordinary file inside the workspace. Creating one
    needs local write access to `documents/` AND read access to the target, which
    is strictly more than an injected document or the worker can reach.

    Raises:
        FileNotFoundError: If file does not exist.
        OSError: If the path cannot be inspected (see is_reparse_point).
        ValueError: If extension doesn't match, path escapes workspace,
                    symlink points outside workspace, or file is too large.
    """
    raw = Path(path)
    # Check BEFORE resolving — resolve() follows the link, making the check False.
    # is_reparse_point() rather than is_symlink(): it also catches an NTFS junction
    # (which is_symlink() reports False for), and it fails CLOSED, where
    # Path.is_symlink() turns every OSError into "not a link" on Python 3.14 — one
    # of the two interpreters CI runs.
    if is_reparse_point(raw):
        _assert_within(raw, workspace, subject="Symlink", verb="points")
    p = _assert_within(raw, workspace, subject="Path", verb="is")
    if not p.is_file():
        raise FileNotFoundError(f"File not found: {redact(path)}")
    if expected_suffixes and p.suffix.lower() not in expected_suffixes:
        raise ValueError(
            f"Unexpected file type '{p.suffix}' for {redact(path)}. "
            f"Expected: {expected_suffixes}")
    file_size = p.stat().st_size
    if file_size > MAX_FILE_SIZE:
        raise ValueError(
            f"File too large: {redact(path)} is {file_size / 1024 / 1024:.1f} MB "
            f"(limit: {MAX_FILE_SIZE / 1024 / 1024:.0f} MB)"
        )
    return p


_WORD_GAP = 3   # pdfplumber's default x_tolerance: glyphs further apart than this begin a new word


_LEVEL = 0.01      # degrees off level at most for a glyph to keep the page's geometry
_FRAME_GAP = 1.0   # degrees apart at most for two rotated glyphs' baselines to read as one run


def _angle(char: dict) -> float:
    """The glyph's baseline direction in degrees, from -180 to 180; level
    for a matrix with no direction or a non-finite one."""
    a, b = char.get("matrix", (1, 0))[:2]
    if not (math.isfinite(a) and math.isfinite(b) and math.hypot(a, b)):
        return 0.0
    return math.degrees(math.atan2(b, a))


def _project(char: dict, angle: float) -> dict:
    """The glyph read along a baseline at `angle`, the source kept for
    overlap checks on the page. An axis-aligned glyph box encloses its own
    rotated advance and font height: solve along the better-conditioned
    axis, including at 90°."""
    a, b, _, _, e, f = char["matrix"]
    scale = math.hypot(a, b)
    across, up = abs(a) / scale, abs(b) / scale
    advance = abs(char["adv"] * scale)
    height = ((char["bottom"] - char["top"] - up * advance) / across
              if across >= up else (char["x1"] - char["x0"] - across * advance) / up)
    u, v = math.cos(math.radians(angle)), math.sin(math.radians(angle))
    x, bottom = u * e + v * f, v * e - u * f
    return {**char, "x0": x, "x1": x + advance, "top": bottom - height,
            "bottom": bottom, "size": height, "_overlay_source": char}


def _overlay_frames(chars: list[dict], tick) -> list[list[dict]]:
    """The page's level glyphs, as they are, then the rest in frames of
    baselines within `_FRAME_GAP` degrees of the frame's first, joined
    across the turn from 180 to -180. A rotated frame's glyphs are read
    along that first baseline, from their PDF matrices, so a rotated SSN is
    one run rather than nine unrelated page rows, and one whose glyphs are
    drawn a hair apart stays one. Only a float's noise counts as level: a
    tilted overlay's delimiters share no baseline with its digits on the
    page, and would split its label."""
    page: list[dict] = []
    frames: list[list[tuple[float, dict]]] = []
    for angle, char in sorted(((_angle(c), c) for c in chars), key=lambda item: item[0]):
        tick()
        if abs(angle) <= _LEVEL:
            page.append(char)
        elif frames and angle - frames[-1][0][0] <= _FRAME_GAP:
            frames[-1].append((angle, char))
        else:
            frames.append([(angle, char)])
    if len(frames) > 1 and frames[0][-1][0] + 360 - frames[-1][0][0] <= _FRAME_GAP:
        frames[0] = frames.pop() + frames[0]
    return [page] + [[_project(char, frame[0][0]) for _, char in frame] for frame in frames]


def _overlay_rows(chars: list[dict], tick) -> list[list[dict]]:
    return [row for frame in _overlay_frames(chars, tick) for row in _word_rows(frame)]


def _spelled_with(digits: list[dict], chars: list[dict], tick) -> set[int]:
    """The ids of the glyphs an identifier spelled "000-00-0000" prints
    between its own digits: any one glyph but a letter, digit or space, in
    the digits' size, on their baseline, between them rather than over
    either, and no further from either than half that size or its own width.
    Left with the letters, an overlay's hyphens, dashes or slashes split its
    label's words. A space stays there, since a label's own spaces share an
    overlay's baseline whenever it is drawn in the label's size; so does a
    label's own slash or apostrophe, which lands on an overlay's digit
    rather than between two."""
    marks = [c for c in chars if len(c["text"]) == 1 and not c["text"].isalnum() and not c["text"].isspace()]
    numerals = {id(d) for d in digits}
    lines: dict[tuple, list[dict]] = {}
    between: list[tuple[tuple, dict]] = []
    for n, frame in enumerate(_overlay_frames(digits + marks, tick)):
        for glyph in frame:
            tick()
            key = (n, round(glyph["size"], 1), round(glyph["bottom"], 1))
            if id(glyph.get("_overlay_source", glyph)) in numerals:
                lines.setdefault(key, []).append(glyph)
            else:
                between.append((key, glyph))
    for line in lines.values():
        line.sort(key=lambda d: d["x0"])
    starts = {key: [d["x0"] for d in line] for key, line in lines.items()}
    out = set()
    for key, glyph in between:
        tick()
        line = lines.get(key, [])
        k = bisect.bisect_left(starts.get(key, []), glyph["x0"])
        reach = max(glyph["x1"] - glyph["x0"], glyph["size"] / 2)
        if (0 < k < len(line) and -0.5 <= glyph["x0"] - line[k - 1]["x1"] <= reach
                and -0.5 <= line[k]["x0"] - glyph["x1"] <= reach):
            out.add(id(glyph.get("_overlay_source", glyph)))
    return out


def _overlaid_digits(chars: list[dict], tick, protected=frozenset()) -> set[int]:
    """The ids of the digits, and the delimiters between them, of an SSN
    drawn over its label's letters, which pdfplumber interleaves into words
    no label pattern reads ("Y0ou0r so0cial"). Labels are read from the
    letters alone; the digits and delimiters are read as words of their own,
    split where pdfplumber splits them, and an SSN is what `identifier_spans`
    finds there, as in any page's text. A word whose glyphs cross a label's
    letters takes every SSN it holds, whole: its end can hang outside the
    label. Any other digits interleaved with a label on its line go too, so
    that it reads whole again: pdfplumber garbles them into its words
    already. So do the words running on from them past it, when together
    they hold nine digits no SSN accounts for, however spaced: the label's
    SSN. Digits that only graze a label from the next line print apart from
    it, read as any page's text is. A word in `protected` (its rounded box)
    is one the positional rules already mask, keeping its last four. The
    source page's characters are never modified. Wrapped headings share the
    positional reader's words; rotated digits are grouped in their own
    baseline frame and checked against the original page's label letters."""
    from pdfplumber.utils import extract_words

    digits = [c for c in chars if c["text"] in "0123456789" and len(c["text"]) == 1]
    if not digits:
        return set()
    delimiters = _spelled_with(digits, chars, tick)
    letters = [c for c in chars if not any(d in "0123456789" for d in c["text"]) and id(c) not in delimiters]
    index: dict[tuple[int, int], list[tuple[dict, int]]] = {}   # each label letter, with its label's number
    indexed: set[int] = set()
    indexed_words: set[int] = set()
    entries = 0
    rows = _word_rows(extract_words(letters, return_chars=True))
    covers = []
    for row in rows:
        tick()
        text, starts = _row_text(row)
        phrase_of: dict[int, tuple[list[dict], int]] = {}   # a word's phrase and where it ends in the row's text
        read = 0
        for phrase in _phrases(row):
            read += len(phrase)
            phrase_of.update((id(w), (phrase, starts[read - 1] + len(row[read - 1]["text"]))) for w in phrase)
        for low, high, _ in identifier_spans(text)[0]:
            tick()
            cover = _covered(row, starts, low, high)
            phrase, phrase_end = phrase_of[id(cover[0])]
            if _LABEL_TAIL.fullmatch(text, high, phrase_end):
                cover = phrase
            covers.append(cover)
    covers.extend(phrase + tail for _, tail, phrase in _wrapped_label_words(rows, tick))
    for label, cover in enumerate(covers):
        tick()
        for word in cover:
            if id(word) in indexed_words:
                continue
            indexed_words.add(id(word))
            for char in word["chars"]:
                tick()
                if id(char) in indexed or not char["text"].isalpha():
                    continue
                indexed.add(id(char))
                for x in _strip_span(char["x0"], char["x1"]):
                    for y in _strip_span(char["top"], char["bottom"]):
                        tick()
                        entries += 1
                        if entries > MAX_OVERLAY_INDEX_ENTRIES:
                            raise ParseBudgetError("Identifier overlap index exceeds its per-page budget; split or repair the PDF.")
                        index.setdefault((x, y), []).append((char, label))
    if not index:
        return set()
    lines = {id(c): n for n, line in enumerate(_word_rows(chars)) for c in line}
    hidden: set[int] = set()
    for row in _overlay_rows(digits + [c for c in chars if id(c) in delimiters], tick):
        tick()
        groups: list[list[dict]] = []   # the row's glyphs in pdfplumber's words
        for glyph in row:
            tick()
            if groups and glyph["x0"] - groups[-1][-1]["x1"] <= _WORD_GAP:
                groups[-1].append(glyph)
            else:
                groups.append([glyph])
        spelled: list[str] = []   # the words' text, delimiters as "-"
        owners: list[dict | None] = []   # the glyph behind each character
        for group in groups:
            if owners:
                spelled.append(" ")
                owners.append(None)
            spelled += [glyph["text"] if glyph["text"].isdigit() else "-" for glyph in group]
            owners += group
        text = "".join(spelled)
        _, values, delimited, _ = identifier_spans(text)
        ssns = [[glyph for glyph in owners[low:high] if glyph is not None]
                for low, high in values + [(low, high) for low, high, _ in delimited if text[low + 3:low + 4] == "-"]]
        accounted = {id(glyph): n for n, run in enumerate(ssns) for glyph in run}   # each glyph's SSN
        unread: dict[int, list[int]] = {}   # the words interleaved with each label's letters, by position
        turned: dict[int, list[int]] = {}   # rotated words crossing them over several of the page's lines
        for at, group in enumerate(groups):
            tick()
            numerals = [glyph for glyph in group if glyph["text"].isdigit()]
            sources = [g.get("_overlay_source", g) for g in group]
            bounds = _box([g.get("_overlay_source", g) for g in numerals or group])
            crossed = _labels_crossed(sources, index, lines, tick)
            if not crossed or tuple(round(bounds[k], 3) for k in ("x0", "x1", "top", "bottom")) in protected:
                continue
            for n in {accounted[id(glyph)] for glyph in group if id(glyph) in accounted}:
                hidden.update(id(glyph.get("_overlay_source", glyph)) for glyph in ssns[n])
            across = "_overlay_source" in group[0] and len({lines[id(g)] for g in sources}) > 1
            for label, interleaved in crossed.items():
                if across:
                    turned.setdefault(label, []).append(at)
                elif interleaved:
                    unread.setdefault(label, []).append(at)
            if across:   # of a word set across the page's lines, only a glyph on a label's own garbles it
                hidden.update(id(g) for g in sources if any(_labels_crossed([g], index, lines, tick).values()))
        # A rotated word goes whole only as part of nine digits running on
        # together, an SSN split, never as a date or code that touches a label.
        for label in sorted(unread.keys() | turned.keys()):
            tick()
            crossing = unread.get(label, [])
            reach = crossing + turned.get(label, [])
            low, high = min(reach), max(reach)   # and the words run on from them, off the label too
            while low and _run_on(groups[low - 1], groups[low]):
                low -= 1
            while high + 1 < len(groups) and _run_on(groups[high], groups[high + 1]):
                high += 1
            numerals = [glyph for group in groups[low:high + 1] for glyph in group
                         if glyph["text"].isdigit() and id(glyph) not in accounted]
            if len(numerals) == 9:   # one label's nine digits, however spaced: its SSN
                crossing = list(range(low, high + 1))
            hidden.update(id(glyph.get("_overlay_source", glyph)) for at in crossing for glyph in groups[at])
    return hidden


def _run_on(left: list[dict], right: list[dict]) -> bool:
    """Whether the right word runs on from the left: no further apart than
    twice its glyphs are tall, as a misplaced SSN's groups print."""
    return right[0]["x0"] - left[-1]["x1"] <= 2 * (right[0]["bottom"] - right[0]["top"])


def _labels_crossed(glyphs: list[dict], index: dict, lines: dict[int, int], tick) -> dict[int, bool]:
    """The numbers of the indexed labels whose letters the glyphs' boxes
    cross, each with whether a crossing glyph shares that letter's line in
    `lines` (the page's characters clustered as pdfplumber clusters them),
    where pdfplumber interleaves the two into one word."""
    out: dict[int, bool] = {}
    for glyph in glyphs:
        tick()
        for x in _strip_span(glyph["x0"], glyph["x1"]):
            for y in _strip_span(glyph["top"], glyph["bottom"]):
                tick()
                for char, label in index.get((x, y), ()):
                    if (glyph["x0"] < char["x1"] and glyph["x1"] > char["x0"]
                            and glyph["top"] < char["bottom"] and glyph["bottom"] > char["top"]):
                        out[label] = out.get(label, False) or lines[id(glyph)] == lines[id(char)]
    return out


def _safe_pdf_page(page, tick):
    """A fresh derived page without the digits of an identifier drawn over
    its label, nor the delimiters between them, so the label reads whole
    again and governs its own box. FilteredPage owns its objects mapping;
    the source page and every non-character object (table rules) stay intact.
    """
    tick()
    hidden = _overlaid_digits(page.chars, tick)
    if not hidden:
        return page
    words = page.extract_words()
    governed = prepare_runs(_governed_runs(words, tick)).spelling
    protected = {tuple(round(box[k], 3) for k in ("x0", "x1", "top", "bottom"))
                 for raw, box in _page_identifiers(words, tick)[1]
                 if "".join(re.findall(r"[0-9]", raw)) in governed}
    if protected:
        hidden = _overlaid_digits(page.chars, tick, protected=protected)
    if not hidden:
        return page
    safe = page.filter(lambda obj: True)
    safe.objects["char"] = [c for c in page.chars if id(c) not in hidden]
    tick()
    return safe


def extract_pdf_text(path: str, pages: list[int] | None = None, workspace: str | None = None) -> dict[int, str]:
    """Extract text from PDF, page by page.

    Args:
        path: Path to PDF file.
        pages: Optional list of 0-indexed page numbers. None = all pages.
        workspace: If provided, path must resolve within this directory.

    Returns:
        Dict mapping page number (0-indexed) to extracted text.

    Raises:
        FileNotFoundError: If file does not exist.
        ValueError: If path escapes workspace or file is too large.
        IndexError: If a requested page number is out of range.
        RuntimeError: If the installed parsing stack is not the locked one.
    """
    # BEFORE the import, for the reason `extract_excel` refuses before its own:
    # a document that is going to be refused must not be opened, and the refusal
    # is about the library that would open it.
    _require_locked_parsers()

    import pdfplumber

    checked = _validate_file_path(path, {".pdf"}, workspace=workspace)

    budget = _Budget(f"Reading {redact(path)}")
    result = {}
    with pdfplumber.open(checked) as pdf:
        num_pages = len(pdf.pages)
        target_pages = pages if pages is not None else range(num_pages)
        _refuse_too_many_pages(len(target_pages), path)
        for i in target_pages:
            if i < 0 or i >= num_pages:
                raise IndexError(
                    f"Page {i} out of range for {redact(path)} "
                    f"({num_pages} pages)")
            budget.tick()
            page = _safe_pdf_page(pdf.pages[i], budget.tick)
            text = page.extract_text() or ""
            budget.tick()   # the page itself may have crossed the deadline
            budget.charge(text)
            runs = _governed_runs(page.extract_words(), budget.tick)
            result[i] = mask_runs(mask_ssn(text), runs)
    budget.tick()
    return result


# pdfplumber's own line tolerance: words whose tops differ by this much share a row.
_ROW_TOLERANCE = 3


def _word_rows(words: list[dict]) -> list[list[dict]]:
    """Rows as pdfplumber's text forms them: each word joins the row of the
    word above it in reading order when their tops are within the tolerance,
    so a row can step down in small changes of font size."""
    rows: list[list[dict]] = []
    previous = None
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if rows and word["top"] - previous <= _ROW_TOLERANCE:
            rows[-1].append(word)
        else:
            rows.append([word])
        previous = word["top"]
    return [sorted(row, key=lambda w: w["x0"]) for row in rows]


# A label governs the value printed beside it on its row or just under it in
# its column: the box a form prints. Page text loses that. A rotated sidebar,
# a form-code row or a margin glyph can land between the label's line and the
# value's (Form 540's header SSN, W-2 and 1099 copies), so every PDF route also
# masks by position. Distances are in the label's own height; the filled
# official forms put a value at most 9.1 heights right of its label, the next
# dependent 14.4 beyond the last (Form 540), 1.8 to one side below it, and a
# Form 2441 provider 12.3 down its column. A dot leader printed right after a
# label carries it across the row (FTB 3532's line 3: 42 heights of dots), so
# its first value is measured from the leader's end. Past its heading's reach
# a column runs from entry to entry, as far as a filer leaves rows without a
# number (Form 1095-C's date-of-birth rows), with no word at all printed
# across it between: a statement's own headings print larger than its labels.
_BESIDE = 12.0   # how far right of its label, on its row, its first value may start
_CHAIN = 20.0    # how far beyond one value on that row the next may start
_UNDER = 2.5     # how far below its label a boxed value may print
_ASIDE = 2.0     # how far outside the label's span that value may start
_COLUMN = 14.0   # how far below its heading label a column's first entry may print
_LINED = 0.35    # how far off its first entry's edges or centre a column's entry may print (IT-201's dependents: 0.02)
_RUN = 64.0      # how far below a column's last entry its next may print (a 1095-C page's 13 rows: 60.4)
_DOTS = 2.0      # how far from its label, and from each other, a leader's dots may print (Schedule 1's first: 1.5)
_BIN = 20.0      # the width of the strips a page's words are indexed in
_STRIPS = int(14400 // _BIN)   # a PDF page is at most 14,400 points wide
_LEADER = re.compile(r"[.\u2026]+")   # a word printed as dots alone
_WRAP = 64       # characters of a heading phrase's end read ("individual taxpayer identification": 34)
_CAPTION = re.compile(r"^(?:\((?:[a-z]|[0-9]+[a-z]?)\)|[a-z]\.|[0-9]+[a-z]?)(?:\s|$)", re.I)
_NUMBERED = re.compile(r"^(?:\([0-9]+[a-z]?\)|[0-9]+[a-z]?)(?:\s|$)", re.I)
_EMPTY_ENTRY = re.compile(r"^(?:N/?A|Applied for|Tax[- ]Exempt|Died)$", re.I)   # Schedule EIC line 2: "Died"
_LABEL_TAIL = re.compile(r"[\s.:)]*(?:or\s+other\s+TIN[\s.:)]*)?", re.I)   # what may end a label's phrase after it (1095-B/C)
_ID_COLUMN = re.compile(r"SSN\s+or\s+ITIN", re.I)   # IT-213's unnumbered child identifier column
_ABUT = 0.05     # points apart at most for words a hidden glyph split (TY2024 1040's "num b e r")


def _row_text(row: list[dict]) -> tuple[str, list[int]]:
    """The row's words joined by a space unless they abut, and where each
    starts: the TY2024 Form 1040 draws spaces over "number" in its
    dependents' heading, splitting it into words that touch."""
    starts, parts, offset = [], [], 0
    for i, word in enumerate(row):
        if i:
            gap = "" if word["x0"] - row[i - 1]["x1"] <= _ABUT else " "
            parts.append(gap)
            offset += len(gap)
        starts.append(offset)
        parts.append(word["text"])
        offset += len(word["text"])
    return "".join(parts), starts


def _phrases(row: list[dict]) -> list[list[dict]]:
    """A row's words in groups no further apart than a word is tall: the
    cells of a heading row."""
    out: list[list[dict]] = []
    for word in row:
        if out and word["x0"] - out[-1][-1]["x1"] <= word["bottom"] - word["top"]:
            out[-1].append(word)
        else:
            out.append([word])
    return out


def _wrapped_label_words(rows: list[list[dict]], tick) -> list:
    """Identifier headings printed on two rows: "(2) Social security" over
    "number" (the TY2024 Form 1040's dependents) and "(2) Identifying" over
    "number" (Form 1040-NR's). A label starts in the upper phrase and ends
    the lower one, which holds nothing else and starts no further below the
    upper than it is tall. Return its matched upper words, lower words,
    and complete upper phrase: overlap detection includes "Your" before
    the match, while the positional reader keeps the match's own box.
    Only an upper phrase's last `_WRAP` characters are read, where such a
    label starts."""
    labels = []
    for upper, lower in zip(rows, rows[1:]):
        tick()
        heads = _phrases(upper)
        ends = [head[-1]["x1"] for head in heads]
        read: dict[int, tuple] = {}
        for tail in _phrases(lower):
            tick()
            j = bisect.bisect_right(ends, tail[0]["x0"])
            if j == len(heads) or heads[j][0]["x0"] >= tail[-1]["x1"]:
                continue
            if j not in read:
                read[j] = (*_row_text(heads[j]), max(w["bottom"] for w in heads[j]))
            head_text, starts, head_bottom = read[j]
            top, bottom = tail[0]["top"], max(w["bottom"] for w in tail)
            if top - head_bottom > bottom - top:
                continue
            cut = max(len(head_text) - _WRAP, 0)
            joined = f"{head_text[cut:]} {_row_text(tail)[0]}"
            low = next((cut + start for start, end, _ in identifier_spans(joined)[0]
                        if start < len(head_text) - cut and end == len(joined)), None)
            if low is not None:
                labels.append((_covered(heads[j], starts, low, len(head_text)), tail, heads[j]))
    return labels


def _wrapped_labels(rows: list[list[dict]], tick) -> list:
    labels = []
    for head, tail, _ in _wrapped_label_words(rows, tick):
        box = {**_box(head + tail), "top": tail[0]["top"], "bottom": max(w["bottom"] for w in tail)}
        labels.append((box, False, box["x1"]))
    return labels


def _covered(row: list[dict], starts: list[int], start: int, end: int) -> list[dict]:
    return row[max(bisect.bisect_right(starts, start) - 1, 0):bisect.bisect_left(starts, end)]


def _box(words: list[dict]) -> dict:
    return {"x0": min(w["x0"] for w in words), "x1": max(w["x1"] for w in words),
            "top": min(w["top"] for w in words), "bottom": max(w["bottom"] for w in words)}


def _leader_end(label: dict, row: list[dict], k: int) -> float:
    """Where the dot leader starting at `row[k]`, right after `label`, ends;
    the label's own end if there is none."""
    end = label["x1"]
    while (k < len(row) and _LEADER.fullmatch(row[k]["text"])
           and row[k]["x0"] - end <= _DOTS * (label["bottom"] - label["top"])):
        end = row[k]["x1"]
        k += 1
    return end


def _page_identifiers(words: list[dict], tick) -> tuple[list, list, list, set]:
    """The page's identifier labels (box, carries, where its row's reach
    starts), runs still to mask (text, box), identifiers spelled another way
    (box, masked by the page itself), and the ids of every word an
    identifier or a label's note covers: none of them ends a column."""
    labels: list[tuple[dict, bool, float]] = []
    values: list[tuple[str, dict]] = []
    spelled: list[tuple[dict, bool]] = []
    covered: set[int] = set()
    rows = _word_rows(words)
    for row in rows:
        tick()
        text, starts = _row_text(row)
        found_labels, found_values, found_spelled, notes = identifier_spans(text)
        row_top = min(w["top"] for w in row)
        phrases = _phrases(row)
        phrase_of: dict[int, tuple[list[dict], str, int]] = {}   # a word's phrase, its text, and where it ends in the row's
        read = 0
        for phrase in phrases:
            tick()
            read += len(phrase)
            phrase_end = starts[read - 1] + len(row[read - 1]["text"])
            phrase_of.update((id(w), (phrase, _row_text(phrase)[0], phrase_end)) for w in phrase)
        leaders: dict[tuple, float] = {}   # labels in one word ("SSN/ITIN") share its leader
        for low, high, carries in found_labels:
            tick()
            cover = _covered(row, starts, low, high)
            box = _box(cover)
            phrase, phrase_text, phrase_end = phrase_of[id(cover[0])]
            ends_phrase = bool(_LABEL_TAIL.fullmatch(text, high, phrase_end))
            box["row_top"] = row_top
            box["phrase_end"] = ends_phrase
            box["phrase_x0"] = phrase[0]["x0"]
            box["numbered"] = bool(_NUMBERED.match(phrase_text)) and phrase is phrases[0] and ends_phrase
            box["caption"] = (bool(_CAPTION.match(phrase_text)) or bool(_ID_COLUMN.fullmatch(phrase_text))) and ends_phrase and box["x0"] - phrase[0]["x0"] <= _BESIDE * (box["bottom"] - box["top"])
            if box["caption"]:
                box["x0"] = phrase[0]["x0"]
            k = bisect.bisect_left(starts, high)
            key = (k, box["x1"], box["bottom"] - box["top"])
            if key not in leaders:
                leaders[key] = _leader_end(box, row, k)
            labels.append((box, carries, leaders[key]))
        for low, high in found_values:
            cover = _covered(row, starts, low, high)
            covered.update(id(w) for w in cover)
            values.append((text[low:high], _box(cover)))
        for low, high, masked in found_spelled:
            cover = _covered(row, starts, low, high)
            covered.update(id(w) for w in cover)
            spelled.append((_box(cover), masked))
        for low, high in notes:
            covered.update(id(w) for w in _covered(row, starts, low, high))
    return labels + _wrapped_labels(rows, tick), values, spelled, covered


def _is_under(box: dict, label: dict, height: float) -> bool:
    """`box` prints just under `label`, in the box the label heads."""
    return (box["x1"] >= label["x0"] - _ASIDE * height and box["x0"] <= label["x1"] + _ASIDE * height
            and -1 <= box["top"] - label["bottom"] <= _UNDER * height)


def _just_under(labels: list, items: list, tick) -> list:
    """For each label, the identifiers (box, run text or None if spelled
    otherwise, masked by the page) printed just under it."""
    items = sorted(items, key=lambda item: item[0]["top"])
    tops = [item[0]["top"] for item in items]
    out = []
    for label, *_ in labels:
        tick()
        height = label["bottom"] - label["top"]
        window = items[bisect.bisect_left(tops, label["bottom"] - 1):bisect.bisect_right(tops, label["bottom"] + _UNDER * height)]
        out.append([item for item in window if _is_under(item[0], label, height)])
    return out


def _beside(labels: list, items: list, under: list, tick) -> set[str]:
    """Each label's values on its row: the first identifier to its right
    within reach, then each one within reach of the one before (Form 540's
    dependents, "Spouse" or a glyph between). A label whose value prints
    under it has none on its row; one whose first value the page masked
    itself shows last fours only, and governs nothing after it (a broker's
    account number further along). A filer's hyphens say nothing about the
    next value."""
    items = sorted(items, key=lambda item: item[0]["top"])
    tops = [item[0]["top"] for item in items]
    tallest = max((item[0]["bottom"] - item[0]["top"] for item in items), default=0.0)
    runs: set[str] = set()
    for (label, _, start), below in zip(labels, under):
        tick()
        height = label["bottom"] - label["top"]
        window = items[bisect.bisect_left(tops, label["top"] - tallest):bisect.bisect_left(tops, label["bottom"])]
        row = sorted((item for item in window if item[0]["bottom"] > label["top"] and item[0]["x0"] >= label["x1"] - 1),
                     key=lambda item: item[0]["x0"])
        if below or not row or row[0][2]:
            continue
        reach, limit = start, _BESIDE
        for box, raw, _ in row:
            if box["x0"] - reach > limit * height:
                break
            if raw is not None:
                runs.add(raw)
            reach, limit = box["x1"], _CHAIN
    return runs


def _strip(x: float) -> int:
    """The vertical strip of the page `x` falls in, the page's first or last
    for an `x` off it."""
    return min(max(int(x // _BIN), 0), _STRIPS)


def _strip_span(x0: float, x1: float) -> range:
    """The vertical strips of the page from `x0` to `x1`."""
    return range(_strip(x0), _strip(x1) + 1)


def _strips(words: list[dict]) -> dict[int, list]:
    """The words indexed by the vertical strips of the page they cross."""
    index: dict[int, list] = {}
    for w in words:
        for strip in _strip_span(w["x0"], w["x1"]):
            index.setdefault(strip, []).append(w)
    return index


def _over(index: dict, box: dict) -> list[dict]:
    """The indexed words printed above `box` and overlapping its span."""
    seen: dict[int, dict] = {}
    for strip in _strip_span(box["x0"], box["x1"]):
        for w in index.get(strip, ()):
            if w["x1"] > box["x0"] and w["x0"] < box["x1"] and w["bottom"] <= box["top"] + 1:
                seen[id(w)] = w
    return list(seen.values())


def _by_height(labels: list) -> list:
    """Labels in classes of height within a factor of two, each sorted by
    bottom, so one tall label widens only its own class's reach."""
    classes: dict[int, list] = {}
    for k, (label, *_) in enumerate(labels):
        height = label["bottom"] - label["top"]
        classes.setdefault(int(math.log2(max(height, 1.0))), []).append((label["bottom"], k, height))
    return [(sorted(members), max(h for _, _, h in members)) for members in classes.values()]


def _in_column(box: dict, above: list[dict], label: dict, height: float, strict: bool = False) -> bool:
    """`box` is in the column `label` heads: below it, in its span, and no
    word printed between, across `box`, as short as the label's own text,
    or when `strict` no word at all."""
    ceiling = next((w["bottom"] for w in above if strict or w["bottom"] - w["top"] <= height + 1), float("-inf"))
    return (box["top"] > label["bottom"] >= ceiling - 1
            and box["x1"] > label["x0"] - height and box["x0"] < label["x1"] + height)


def _edges(box: dict) -> tuple[float, float, float]:
    """The x of `box`'s left edge, right edge and centre: a column's entries
    line up on one, by how the form's boxes align what is typed in them."""
    return box["x0"], box["x1"], (box["x0"] + box["x1"]) / 2


def _lines_up(box: dict, column: tuple) -> bool:
    """`box` prints with an edge or its centre where `column`'s first entry does."""
    _, height, *first = column
    return any(abs(x - at) <= _LINED * height for x, at in zip(_edges(box), first))


def _edge_strips(column: tuple) -> set[int]:
    """The strips an entry's edge or centre falls in when it lines up with
    `column`'s first entry."""
    _, height, *first = column
    reach = _LINED * height
    return {strip for x in first for strip in _strip_span(x - reach, x + reach)}


def _lined_columns(columns: dict, lined: dict, box: dict) -> list:
    """The columns `box` may line up with, as (label index, height, column):
    those filed under the strips its edges and centre fall in, within reach
    of their last entry. Entries come top to bottom, so a column out of
    reach of one is out of reach of every later one, and is dropped."""
    live = []
    for k in sorted(set().union(*(lined.get(_strip(x), ()) for x in _edges(box)))):
        last, height = columns[k][:2]
        if box["top"] - last <= _RUN * height:
            live.append((k, height, columns[k]))
        else:
            _drop_column(columns, lined, k)
    return live


def _drop_column(columns: dict, lined: dict, k: int) -> None:
    for strip in _edge_strips(columns.pop(k)):
        lined[strip].discard(k)


def _start_column(columns: dict, lined: dict, k: int, box: dict, height: float) -> None:
    """`box` is the first entry of label `k`'s column: file the column under
    the strips a lined-up entry's edges fall in, taking over each column the
    entry lines up with. A heading lower down starts a table of its own, and
    a strip holds no more columns than it has places to line up."""
    for j, _, column in _lined_columns(columns, lined, box):
        if _lines_up(box, column):
            _drop_column(columns, lined, j)
    columns[k] = (box["bottom"], height, *_edges(box))
    for strip in _edge_strips(columns[k]):
        lined.setdefault(strip, set()).add(k)


def _under(labels: list, items: list, under: list, printed: list[dict], tick) -> set[str]:
    """Runs under a label with no value on its row: just below it, or down
    its column with nothing printed between but other identifiers, notes on
    the label, and the entries typed in the column's boxes, which print
    larger than the form's own labels (Form 2441's providers, one of them
    "Tax-Exempt"). Past the label's own reach a column runs on from entry
    to entry, each lined up with its first by an edge or the centre, within
    `_RUN` heights of the last, and with no word at all printed across it
    between: past as many rows as a filer leaves without a number (Form
    1095-C's date-of-birth rows), whatever the filer's hyphens do to an
    entry's width in a centred box (Form 1095-B's), but never past a
    statement's own heading, printed larger than its labels. An entry the
    page masked itself neither starts a column nor ends one, and a label whose own value just
    under it the page masked is a box's label on a statement, heading no
    column; a filer's hyphens in the first entry do not end it (Form 2441's
    qualifying persons). A header SSN typed in a comb box spans it, so its
    right edge can line up with the amounts down the margin under it (Form
    5695's page 2), and the column masks those that are identifier-shaped;
    IT-201's page 3 header is 3 points off its amounts."""
    free = [k for k, (_, carries, _) in enumerate(labels) if not carries]
    heads = {k for k in free if not any(masked for _, _, masked in under[k])}
    classes = [(reach, members, tallest, group) for caption, reach in ((False, _COLUMN), (True, _RUN))
               for group in [[k for k in free if bool(labels[k][0].get("caption")) == caption]]
               for members, tallest in _by_height([labels[k] for k in group])]
    index = _strips(printed)
    columns: dict[int, tuple] = {}     # a column's last entry's bottom, label height, first entry's edges and centre
    lined: dict[int, set[int]] = {}    # the columns an entry with an edge or centre in each strip may line up with
    runs = set()
    for box, raw, masked in sorted(items, key=lambda item: item[0]["top"]):
        tick()
        if masked:
            continue
        above = sorted(_over(index, box), key=lambda w: -w["bottom"])
        window = [(group[j], height, None) for reach, members, tallest, group in classes
                  for _, j, height in members[bisect.bisect_left(members, (box["top"] - reach * tallest,)):
                                              bisect.bisect_right(members, (box["top"] + 1, float("inf")))]]
        for k, height, column in window + _lined_columns(columns, lined, box):
            tick()
            label = labels[k][0]
            if column is not None:
                governed = _lines_up(box, column) and _in_column(box, above, label, height, strict=True)
            else:   # a spelled entry is masked already, and counts only toward a column it starts
                governed = ((raw is not None and _is_under(box, label, height))
                            or (k in heads and box["top"] - label["bottom"] <= (_RUN if label.get("caption") else _COLUMN) * height
                                and _in_column(box, above, label, height, strict=box["top"] - label["bottom"] > _COLUMN * height)))
            if governed:
                if raw is not None:
                    runs.add(raw)
                if k in columns:
                    columns[k] = (box["bottom"], *columns[k][1:])
                elif k in heads:
                    _start_column(columns, lined, k, box, height)
                break
    return runs


def _lone_labels(labels: list, items: list, under: list, tick) -> list:
    """Free labels with no identifier below or beside, including raised values."""
    out = []
    ordered = sorted(items, key=lambda item: item[0]["top"])
    tops = [item[0]["top"] for item in ordered]
    tallest = max((b["bottom"] - b["top"] for b, *_ in items), default=0.0)
    for entry, below in zip(labels, under):
        tick()
        label, carries, _ = entry
        near = ordered[bisect.bisect_left(tops, label["top"] - tallest):bisect.bisect_left(tops, label["bottom"])]
        if not carries and not below and not any(b["bottom"] > label["top"] and b["x0"] >= label["x1"] - 1 for b, *_ in near):
            out.append(entry)
    return out


def _above(labels: list, items: list, tick) -> set[str]:
    """A signature line's number just above its otherwise empty label, typed
    anywhere in the line's field (Form 8332's is 12 points tall, its label's
    7). Whatever else ends there is masked too, the safe direction: a ZIP+4
    over a 1099's empty TIN, a FEIN typed without its hyphen over FTB 593's
    empty individual row."""
    out = set()
    ordered = sorted(items, key=lambda item: item[0]["bottom"])
    bottoms = [item[0]["bottom"] for item in ordered]
    for label, *_ in labels:
        tick()
        if not label.get("phrase_end"):
            continue
        height = label["bottom"] - label["top"]
        for box, raw, _ in ordered[bisect.bisect_left(bottoms, label["top"] - 1.5 * height):bisect.bisect_right(bottoms, label["top"])]:
            tick()
            if raw is not None and box["x1"] > label.get("phrase_x0", label["x0"]) and box["x0"] < label["x1"]:
                out.add(raw)
    return out


def _paragraph_labels(words: list[dict], labels: list, tick, boundaries: list | None = None) -> list:
    """A numbered identifier question lends only its beside reach to the
    indented, flush instruction lines below it, stopping at the next item."""
    rows = _word_rows(words)
    tops = [min(w["top"] for w in row) for row in rows]
    out = []
    numbered_rows = {label.get("row_top", label["top"]) for label, *_ in (labels if boundaries is None else boundaries) if label.get("numbered")}
    for label, *_ in labels:
        tick()
        if not label.get("numbered"):
            continue
        height = label["bottom"] - label["top"]
        last, left = label["bottom"], None
        for j in range(bisect.bisect_right(tops, label["top"]), len(rows)):
            tick()
            if tops[j] in numbered_rows:
                break
            phrase = _phrases(rows[j])[0]
            box = _box(phrase)
            x = box["x0"]
            if box["top"] - last > height or (left is None and not 0 <= x - label["phrase_x0"] <= 2 * height) or (left is not None and abs(x - left) > 1):
                break
            text = _row_text(phrase)[0]
            _, values, spelled, _ = identifier_spans(text)
            if values or spelled:
                break
            out.append((box, False, box["x1"]))
            left, last = x, box["bottom"]
    return out


def _governed_runs(words: list[dict], tick=lambda: None) -> set[str]:
    """The text of each identifier-shaped run the page prints in an
    identifier label's box, whatever order the page's text puts them in.
    `tick` is the caller's time budget, checked as the page is read."""
    labels, values, spelled, covered = _page_identifiers(words, tick)
    if not labels or not values:
        return set()
    items = [(box, raw, False) for raw, box in values] + [(box, None, masked) for box, masked in spelled]
    under = _just_under(labels, items, tick)
    placeholders = []   # typed in an empty box: no column ends at one, and a row's reach steps over it
    for row in _word_rows(words):
        tick()
        for phrase in _phrases(row):
            tick()
            if _EMPTY_ENTRY.fullmatch(_row_text(phrase)[0]):
                covered.update(id(w) for w in phrase)
                placeholders.append((_box(phrase), None, False))
    printed = [w for w in words if id(w) not in covered]
    lone = _lone_labels(labels, items, under, tick)
    extra = _paragraph_labels(words, lone, tick, boundaries=labels)
    return (_beside(labels + extra, items + placeholders, under + [[] for _ in extra], tick)
            | _under(labels, items, under, printed, tick) | _above(lone, items, tick))


def _mask_table_runs(table: list, runs, tick) -> list:
    """The runs a page's labels govern, masked in each cell of a table and
    across each row's cells."""
    out = []
    for row in table:
        tick()
        if isinstance(row, list):
            row = mask_row_runs([mask_runs(cell, runs) if isinstance(cell, str) else cell for cell in row], runs)
        out.append(row)
    return out


def _identity_identifiers(row: list[dict], safe: str) -> list[dict]:
    """Positions of identifiers already masked in this row; never raw digits."""
    text = ' '.join(w['text'] for w in row)
    starts, offset = [], 0
    for word in row:
        starts.append(offset)
        offset += len(word['text']) + 1
    _, values, spelled, _ = identifier_spans(text)
    allowed = set(re.findall(r'(?<![A-Za-z0-9])XXX-XX-[0-9]{4}(?![A-Za-z0-9])', safe))
    out = {}
    for low, high in values + [(a, b) for a, b, _ in spelled]:
        positions = [low + m.start() for m in re.finditer(r'[0-9]', text[low:high])]
        if len(positions) == 4 and text[low:high] in allowed:
            chunks = [(low, high, text[low:high])]
        elif positions and len(positions) % 9 == 0:
            chunks = [(positions[i], positions[i + 8] + 1,
                       'XXX-XX-' + ''.join(text[j] for j in positions[i + 5:i + 9]))
                      for i in range(0, len(positions), 9)]
        else:
            continue
        for left, right, masked in chunks:
            if masked not in allowed:
                continue
            bounds = _box(_covered(row, starts, left, right))
            key = (masked, *(round(float(bounds[k]), 1) for k in ('x0', 'x1', 'top', 'bottom')))
            out[key] = dict(zip(('text', 'x0', 'x1', 'top', 'bottom'), key))
    return list(out.values())


def _masked_words(words: list[dict], tick=lambda: None, *, identity_layout: bool = False) -> list[dict]:
    """`mask_ssn` over the page's rows, top to bottom, so a label governs the
    row printed under it whatever order the content stream drew them in, then
    the runs a label governs by position. A row the masker changes comes back
    as ONE word spanning it, holding the masked row: no position inside it
    can come back without the identifier."""
    rows = _word_rows(words)
    texts = [" ".join(word["text"] for word in row) for row in rows]
    masked = mask_runs(mask_ssn("\n".join(texts)), _governed_runs(words, tick)).split("\n")
    if len(masked) != len(rows):
        raise RuntimeError("masking changed the page's row count")
    out = []
    for row, text, safe in zip(rows, texts, masked):
        if safe == text:
            out += [{"text": w["text"], **{k: round(float(w[k]), 1) for k in ("x0", "x1", "top", "bottom")}}
                    for w in row]
        else:
            box = {"text": safe, "x0": round(float(min(w["x0"] for w in row)), 1),
                   "x1": round(float(max(w["x1"] for w in row)), 1),
                   "top": round(float(min(w["top"] for w in row)), 1),
                   "bottom": round(float(max(w["bottom"] for w in row)), 1)}
            if identity_layout:
                # A consumer may locate names, but gets no original numeric
                # token from a changed row. Masked identifiers remain together
                # in the aggregate row; no component can reconstruct a full ID.
                box["masked_row"] = True
                box["identity_identifiers"] = _identity_identifiers(row, safe)
                out += [{"text": w["text"], **{k: round(float(w[k]), 1) for k in ("x0", "x1", "top", "bottom")}}
                        for w in row if not any(c.isdigit() for c in w["text"])]
            out.append(box)
    return out


def extract_pdf_words(path: str, pages: list[int] | None = None, workspace: str | None = None, *,
                      identity_layout: bool = False) -> dict[int, list[dict]]:
    """Extract each page's words with their positions, for a reader that finds
    a value by where the form prints it: the row of a line number and the
    column under a heading, which text order does not record.

    Returns a dict mapping page number (0-indexed) to words, each
    `{"text", "x0", "x1", "top", "bottom"}` in points from the page's top left.
    With `identity_layout=True`, changed rows also carry safe nonnumeric word
    positions and masked identifier boxes; original numeric tokens stay hidden.
    Raises as `extract_pdf_text` does.
    """
    _require_locked_parsers()

    import pdfplumber

    checked = _validate_file_path(path, {".pdf"}, workspace=workspace)

    budget = _Budget(f"Reading the words in {redact(path)}")
    result = {}
    with pdfplumber.open(checked) as pdf:
        num_pages = len(pdf.pages)
        target_pages = pages if pages is not None else range(num_pages)
        _refuse_too_many_pages(len(target_pages), path)
        for i in target_pages:
            if i < 0 or i >= num_pages:
                raise IndexError(
                    f"Page {i} out of range for {redact(path)} "
                    f"({num_pages} pages)")
            budget.tick()
            page = _safe_pdf_page(pdf.pages[i], budget.tick)
            words = page.extract_words()
            budget.tick()
            for word in words:
                budget.charge_cell(word["text"])
            result[i] = _masked_words(words, budget.tick, identity_layout=identity_layout)
            if identity_layout:
                for word in result[i]:
                    budget.charge_cell(word["text"])
                    for identifier in word.get("identity_identifiers", []):
                        budget.charge_cell(identifier["text"])
    budget.tick()
    return result


def extract_tables(path: str, pages: list[int] | None = None, workspace: str | None = None) -> dict[int, list[list]]:
    """Extract tables from PDF, page by page.

    Args:
        path: Path to PDF file.
        pages: Optional list of 0-indexed page numbers. None = all pages.
        workspace: If provided, path must resolve within this directory.

    Returns:
        Dict mapping page number to list of tables (each table is list of rows).

    Raises:
        FileNotFoundError: If file does not exist.
        ValueError: If path escapes workspace or file is too large.
        IndexError: If a requested page number is out of range.
        RuntimeError: If the installed parsing stack is not the locked one.
    """
    _require_locked_parsers()

    import pdfplumber

    checked = _validate_file_path(path, {".pdf"}, workspace=workspace)

    budget = _Budget(f"Reading the tables in {redact(path)}")
    result = {}
    with pdfplumber.open(checked) as pdf:
        num_pages = len(pdf.pages)
        target_pages = pages if pages is not None else range(num_pages)
        _refuse_too_many_pages(len(target_pages), path)
        for i in target_pages:
            if i < 0 or i >= num_pages:
                raise IndexError(
                    f"Page {i} out of range for {redact(path)} "
                    f"({num_pages} pages)")
            budget.tick()
            page = _safe_pdf_page(pdf.pages[i], budget.tick)
            tables = page.extract_tables() or []
            # AFTER the extraction, and outside the `if`: a page that yields no
            # table is the one that charges nothing, and the last such page had
            # no check left after it at all.
            budget.tick()
            if tables:
                _charge_deep(budget, tables)
                runs = prepare_runs(_governed_runs(page.extract_words(), budget.tick))
                result[i] = [_mask_table_runs(mask_table(table), runs, budget.tick) for table in tables]
    budget.tick()
    return result


def extract_excel(path: str, sheet_names: list[str] | None = None, workspace: str | None = None) -> dict[str, list[list]]:
    """Extract data from Excel workbook.

    Args:
        path: Path to .xlsx file.
        sheet_names: Optional list of sheet names. None = all sheets.
        workspace: If provided, path must resolve within this directory.

    Returns:
        Dict mapping sheet name to list of rows (each row is list of cell values).

    Raises:
        FileNotFoundError: If file does not exist.
        ValueError: If path escapes workspace or file is too large.
        RuntimeError: If defusedxml is absent, so there is no XXE guard, or if
            the installed parsing stack is not the locked one.
    """
    # BEFORE the path validation and before openpyxl is imported: a workbook
    # that is going to be refused must not be opened, and `load_workbook` reads
    # XML the moment it is called. See the import guard at the top of this file
    # for why this is a refusal and not a warning.
    if not _XXE_GUARD_ACTIVE:
        raise RuntimeError(
            "defusedxml is not installed, so openpyxl would parse this "
            "workbook with the stdlib XML parser and no XXE guard. Source "
            "documents are untrusted input, so this is a refusal, not a "
            "warning. Install the pinned tree: "
            "sh agents/lib/run.sh -m pip install -r requirements.txt")

    # AFTER the guard check, which is about this same tree and says the sharper
    # thing: an absent `defusedxml` is drift too — the lock pins it — and the
    # message naming the XXE hole is the one worth printing over the message
    # naming a version mismatch.
    _require_locked_parsers()

    import openpyxl

    checked = _validate_file_path(path, {".xlsx"}, workspace=workspace)
    # BEFORE `load_workbook`, which decompresses the whole archive into memory:
    # the 100 MB cap bounds the ZIP, not the XML behind it.
    _refuse_a_zip_bomb(str(checked))

    budget = _Budget(f"Reading {redact(path)}")
    wb = openpyxl.load_workbook(checked, data_only=True)
    try:
        targets = sheet_names if sheet_names is not None else wb.sheetnames
        # A sheet that is not there is a refusal naming the ones that are. It
        # used to be skipped, and `{}` with exit 0 is the one answer that reads
        # as "nothing in this workbook" for a misspelled or wrong-case name —
        # the worker records no figures and the tie-out gate cannot see the
        # balances it never learned about.
        for name in targets:
            if name not in wb.sheetnames:
                raise ValueError(
                    f"No sheet named {redact(name)!r} in {redact(path)}. "
                    f"Sheets: {[redact(n) for n in wb.sheetnames]}")
        # The sheet NAME is a string this module returns, and it was the one
        # that left unmasked: `excel` prints it as a JSON key, so a workbook
        # with a sheet named after an identifier disclosed it on the success
        # path while the refusal above redacted the very same name. The masked
        # key is not a name this CLI can be handed back — no more than the
        # redacted ones in that refusal are — and renaming the sheet is the
        # answer to that; two sheets whose names differ only in the digits
        # being masked would collide into one key, which would drop a sheet of
        # a return, so that is a refusal rather than a silently shorter dict.
        extracted = {}
        for name in targets:
            safe = mask_ssn(name)
            if safe in extracted:
                raise ValueError(
                    f"Two sheets in {redact(path)} mask to the same name "
                    f"{safe!r}, so one would be dropped. Rename them.")
            rows = []
            for row in wb[name].iter_rows(values_only=True):
                budget.tick()
                _charge_deep(budget, row)
                rows.append(row)
            extracted[safe] = mask_table(rows)
        budget.tick()
        return extracted
    finally:
        wb.close()


def extract_csv(path: str, encoding: str = "utf-8-sig", workspace: str | None = None) -> list[list[str]]:
    """Extract data from CSV file.

    Args:
        path: Path to CSV file.
        encoding: File encoding (default utf-8-sig handles BOM).
        workspace: If provided, path must resolve within this directory.

    Returns:
        List of rows (each row is list of string values).

    Raises:
        FileNotFoundError: If file does not exist.
        ValueError: If path escapes workspace or file is too large.
    """
    checked = _validate_file_path(path, {".csv"}, workspace=workspace)

    budget = _Budget(f"Reading {redact(path)}")
    with open(checked, newline="", encoding=encoding) as f:
        rows = []
        for row in csv.reader(f):
            budget.tick()
            _charge_deep(budget, row)
            rows.append(row)
        # Masked BEFORE the last check, not after it: masking is the other
        # substantial pass over the file, and a route that ticks and then does
        # work has left that work outside its own budget. The PDF and workbook
        # routes mask per page and per sheet, inside the loop, so their trailing
        # tick already covers it.
        masked = mask_table(rows)
        budget.tick()
        return masked


def pdf_page_count(path: str, workspace: str | None = None) -> int:
    """Return number of pages in a PDF.

    Raises:
        FileNotFoundError: If file does not exist.
        ValueError: If path escapes workspace or file is too large.
        RuntimeError: If the installed parsing stack is not the locked one.
    """
    # Refused like the two extractors, not exempted as a "mere" count. It is
    # /tax-parse step 2.5's page-count probe, the one that decides whether a
    # document splits into sub-tasks, and it opens the PDF with the same
    # pdfminer stack — so a wrong answer here silently drops pages from the
    # parse that follows.
    _require_locked_parsers()

    import pdfplumber

    checked = _validate_file_path(path, {".pdf"}, workspace=workspace)

    with pdfplumber.open(checked) as pdf:
        return len(pdf.pages)


def classify_document(text: str) -> str:
    """Classify a tax document based on its text content.

    Args:
        text: Extracted text from document (typically first page).
              If None or empty, returns "unknown".

    Returns one of: w2, 1099-int, 1099-div, 1099-b, 1099-r, 1099-sa,
    1099-nec, 1099-misc, 1099-k, ca-592b, ca-593, 3921, 3922, brokerage_statement, foreign_bank,
    tax_return, unknown.
    """
    if not text:
        return "unknown"

    text_upper = text.upper()

    # These statements can refer to one another. Prefer the statement's own
    # title over a form number appearing in its instructions or attachments.
    ca_statement_titles = {
        "ca-592b": "RESIDENT AND NONRESIDENT WITHHOLDING TAX STATEMENT",
        "ca-593": "REAL ESTATE WITHHOLDING STATEMENT",
    }
    for doc_type, title in ca_statement_titles.items():
        if title in text_upper:
            return doc_type

    patterns = {
        # Forms 3921/3922 FIRST: their keywords are maximally specific, while
        # broker-issued 3921/3922 packets routinely mention generic phrases like
        # "cost basis" (1099-b) or "Form 1040" (tax_return) in their cover
        # text — insertion order is the tiebreak.
        "3921": ["EXERCISE OF AN INCENTIVE STOCK OPTION", "FORM 3921"],
        # Form 3922 = ESPP purchase-info (§423). Same shadowing risk as 3921:
        # broker packets say "COST BASIS", so 3922 must precede 1099-b.
        "3922": ["EMPLOYEE STOCK PURCHASE PLAN UNDER SECTION 423",
                 "TRANSFER OF STOCK ACQUIRED THROUGH AN EMPLOYEE STOCK PURCHASE",
                 "FORM 3922"],
        "w2": ["WAGE AND TAX STATEMENT", "FORM W-2"],
        "ca-592b": ["FORM 592-B"],
        "ca-593": ["FORM 593"],
        "1099-int": ["INTEREST INCOME", "FORM 1099-INT"],
        "1099-div": ["DIVIDENDS AND DISTRIBUTIONS", "FORM 1099-DIV"],
        "1099-b": ["PROCEEDS FROM BROKER", "FORM 1099-B", "COST BASIS"],
        "1099-r": ["DISTRIBUTIONS FROM PENSIONS", "FORM 1099-R"],
        "1099-sa": ["DISTRIBUTIONS FROM AN HSA", "FORM 1099-SA"],
        "1099-nec": ["NONEMPLOYEE COMPENSATION", "FORM 1099-NEC"],
        "1099-misc": ["MISCELLANEOUS INFORMATION", "FORM 1099-MISC"],
        "1099-k": ["PAYMENT CARD", "FORM 1099-K"],
        "tax_return": ["FORM 1040", "U.S. INDIVIDUAL INCOME TAX RETURN"],
    }

    for doc_type, keywords in patterns.items():
        if any(kw in text_upper for kw in keywords):
            return doc_type

    # Check for brokerage statement patterns
    brokerage_kw = ["BROKERAGE", "CONSOLIDATED", "TAX REPORTING", "YEAR-END"]
    if any(kw in text_upper for kw in brokerage_kw):
        return "brokerage_statement"

    # Check for non-English (likely foreign bank document)
    ascii_ratio = sum(1 for c in text if ord(c) < 128) / max(len(text), 1)
    if ascii_ratio < 0.7:
        return "foreign_bank"

    return "unknown"


def scan_workspace(workspace_dir: str) -> list[dict]:
    """Scan workspace for parseable documents (non-recursive).

    Scans the top-level directory only. For nested documents, point to the
    specific subdirectory.

    Returns list of dicts with: path, filename, extension, size_kb, doc_type (if PDF).

    Raises:
        NotADirectoryError: If workspace_dir is not a valid directory.
        RuntimeError: If the workspace holds a PDF and the installed parsing
            stack is not the locked one.
    """
    workspace = Path(workspace_dir)
    if not workspace.is_dir():
        raise NotADirectoryError(f"Not a directory: {redact(workspace_dir)}")

    extensions = {".pdf", ".xlsx", ".csv"}
    # Skip known output files
    skip_patterns = {"Tax_Return_", "Summary"}

    documents = []
    for f in sorted(workspace.iterdir()):
        # is_reparse_point(), not is_symlink(): a junction is invisible to the
        # latter, and an lstat failure there reads as "not a link" rather than
        # propagating. It fails closed, so an uninspectable entry raises instead
        # of being silently dropped from the document list — a dropped document
        # understates income, which overstates the refund the user acts on.
        if is_reparse_point(f):
            print(f"WARNING: Skipping symlink {redact(f.name)} — copy the "
                  f"file into the workspace instead.",
                  file=sys.stderr)
            continue
        if not f.is_file() or f.suffix.lower() not in extensions:
            continue
        if any(p in f.name for p in skip_patterns):
            continue

        # A file named after an SSN discloses it — `scan` prints this list
        # verbatim — so the name never reaches the transcript. It used to be
        # MASKED, which kept the digits out and handed the worker a `path`
        # that opens nothing: `scan` is the routing table /tax-parse dispatches
        # from, so the next step was a FileNotFoundError on a document that is
        # sitting right there, and two files can mask to one path. The name is
        # the thing that has to change, so this refuses by name the way the
        # symlink arm above does and says so. Ignoring the warning does not
        # quietly drop the document: tie_out's coverage check reads
        # documents/ itself and fails on a file with no ledger record.
        if names_an_identifier(f.name):
            print(f"WARNING: Skipping {redact(f.name)} — its FILENAME carries "
                  f"an identifier, which this list would disclose. Rename the "
                  f"file (the last four digits are all a name may keep) and "
                  f"scan again.", file=sys.stderr)
            continue

        doc = {
            "path": str(f),
            "filename": f.name,
            "extension": f.suffix.lower(),
            "size_kb": round(f.stat().st_size / 1024, 1),
        }

        # Classify by extension
        if f.suffix.lower() == ".pdf":
            # OUTSIDE the try, and per-PDF rather than at the top of the scan.
            # Outside, because that handler turns every exception into a
            # `doc_type: "error"` entry and a stderr line — a refusal routed
            # through it degrades into a warning, which is the exact shape the
            # defusedxml guard above was rewritten to stop being. Per-PDF,
            # because a workspace holding only CSVs opens no locked parser and
            # /tax-parse-crypto works on exactly that.
            #
            # This is the entry point most easily missed: `scan documents/` is
            # the FIRST command /tax-parse runs, it opens every PDF to classify
            # it, and its output is the routing table for everything after — a
            # document classified by an untested parser is one parsed as the
            # wrong form, or dropped.
            _require_locked_parsers()
            try:
                import pdfplumber
                with pdfplumber.open(str(f)) as pdf:
                    first_page_text = pdf.pages[0].extract_text() or "" if pdf.pages else ""
                    doc["doc_type"] = classify_document(first_page_text)
                    doc["page_count"] = len(pdf.pages)
            except ImportError:
                doc["doc_type"] = "error"
                doc["error"] = "pdfplumber not installed"
            except Exception as e:
                doc["doc_type"] = "error"
                # Sanitize error message — avoid leaking document content
                # that may contain SSNs or account numbers. Truncation was the
                # only sanitizing here, and a 200-character prefix of a parser
                # error is where an SSN would sit: redacted at ASSIGNMENT so
                # both consumers are covered, the stderr warning below and the
                # `error` field that rides out in the document list.
                err_msg = redact(e)
                if len(err_msg) > 200:
                    err_msg = err_msg[:200] + "..."
                doc["error"] = f"{type(e).__name__}: {err_msg}"
                print(f"WARNING: Failed to classify {redact(f.name)}: "
                      f"{type(e).__name__}: {redact(err_msg)}",
                      file=sys.stderr)
        elif f.suffix.lower() == ".xlsx":
            doc["doc_type"] = "excel"
        elif f.suffix.lower() == ".csv":
            doc["doc_type"] = "csv"

        documents.append(doc)

    return documents


if __name__ == "__main__":
    # CLI: python parse_utils.py scan <dir>
    # CLI: python parse_utils.py classify <pdf>
    # CLI: python parse_utils.py extract <pdf> [page_num]
    # CLI: python parse_utils.py tables <pdf> [page_num]
    # CLI: python parse_utils.py csv <csv>
    # CLI: python parse_utils.py excel <xlsx> [sheet_name]
    #
    # `csv` and `excel` exist so that a workbook or a CSV has a masked route
    # into the worker's transcript at all. Without them the parse-task template
    # named `extract` for every file, which refuses anything but a PDF, and the
    # worker's only way to read the Excel/CSV group was its own file-read tool
    # or an ad-hoc interpreter call — both of which bypass `mask_ssn` and one of
    # which rule 5 of that template forbids.
    #
    # EVERY READ IS CONFINED TO A WORKSPACE, and the CLI is where that was
    # missing. The library functions have taken a `workspace=` since they were
    # written and refuse a path — or a symlink target — outside it; the CLI
    # passed None to all of them, so the confinement was off for the only caller
    # that matters. That made this command an ARBITRARY FILE READER for anything
    # with a .pdf/.csv/.xlsx suffix: `csv ~/other-return/documents/bank.csv`
    # printed another taxpayer's statement, and a worker allowed to run "only
    # parse_utils" was allowed to read the whole disk one suffix at a time.
    #
    # The default is the CURRENT DIRECTORY, which is the workspace root every
    # skill invokes this from (`run.sh` computes its own root but does not `cd`,
    # so the caller's cwd survives). That makes every documented call confined
    # without changing it. It is deliberately NOT `run.sh`'s derived root: the
    # workspace is where the return lives and the clone is where the code lives,
    # and Y4 is the open item to stop assuming those are the same directory.
    #
    # THE PRECISE CLAIM, because the loose one was written here first and was
    # false: there is no way to run this command with no confinement at all —
    # every route resolves against a workspace. What the flag does is CHOOSE
    # that workspace, so a caller who passes `/` has chosen a wide one. That is
    # a decision available to whoever composes the command line, and in ENFORCED
    # mode the parse worker composes nothing, holding no shell. In SEQUENTIAL
    # mode it is available to a parser that already has a shell, so this is not
    # the layer that stops it — the prompt rules and the banner are.
    # The flag is accepted FIRST or not at all, and a later one is a REFUSAL
    # rather than a wider workspace. Scanning the whole argv for it meant
    # `csv doc.csv --workspace /` read the file: the confinement was still on,
    # applied to `/`, which is the off switch this block says it does not have.
    # Position is what makes that unwritable — a trailing flag now errors, and a
    # caller who meant it is told rather than silently confined to the cwd.
    argv = sys.argv[1:]
    # `$TAX_WORKSPACE` when the flag is absent, and the current directory when
    # neither is set — the same precedence every other CLI here follows, spelled
    # differently only because this one parses its own argv (the command word
    # comes first and argparse cannot express that).
    workspace = None
    if argv and argv[0].startswith("--workspace="):
        workspace = argv[0].split("=", 1)[1]
        del argv[0]
        if not workspace:
            print("Error: --workspace needs a directory", file=sys.stderr)
            sys.exit(1)
    elif argv and argv[0] == "--workspace":
        if len(argv) < 2:
            print("Error: --workspace needs a directory", file=sys.stderr)
            sys.exit(1)
        workspace = argv[1]
        del argv[0:2]
    if any(a == "--workspace" or a.startswith("--workspace=") for a in argv):
        print("Error: --workspace must come first, before the command",
              file=sys.stderr)
        sys.exit(1)
    try:
        workspace = str(resolve_workspace(workspace))
    except WorkspaceError as e:
        print(f"Error: {redact(e)}", file=sys.stderr)
        sys.exit(1)

    if len(argv) < 2:
        print("Usage: python parse_utils.py [--workspace DIR] "
              "<scan|classify|extract|tables|csv|excel> <path> [page|sheet]")
        sys.exit(1)

    cmd, path = argv[0], argv[1]
    # A RELATIVE path means a file in the WORKSPACE. `documents/w2.pdf` is how
    # every skill spells it, and that used to be the same directory the command
    # ran from — which under $TAX_WORKSPACE is the checkout (Y4). Resolving the
    # workspace above and leaving this alone was worse than doing neither: the
    # containment check below then refused the checkout's own `documents/` for
    # being outside the workspace, so `scan documents/` failed rather than
    # scanning the taxpayer's. An ABSOLUTE path is left as given and the
    # containment below is what judges it, exactly as before.
    if not Path(path).is_absolute():
        path = str(Path(workspace) / path)
    arg3 = argv[2] if len(argv) > 2 else None

    def _confined_dir(candidate: str, base: str) -> str:
        """`scan`'s directory, held to the same containment as every file read.

        `scan_workspace` takes a DIRECTORY and lists what is parseable in it, so
        `_validate_file_path` never sees it — and an unconfined `scan /` is a
        map of every tax document on the machine, which is the reconnaissance
        step before the reads this confinement exists to stop.
        """
        if is_reparse_point(Path(candidate)):
            _assert_within(candidate, base, subject="Symlink", verb="points")
        return str(_assert_within(candidate, base, subject="Path", verb="is"))

    def _print_json(obj):
        # `ensure_ascii=False` because the worker reads this: a Cyrillic
        # column header as `\u0414\u0430...` is a decode-by-hand exercise per
        # cell, and the stdout run.sh hands us is UTF-8.
        print(json.dumps(obj, indent=2, ensure_ascii=False, default=str))

    try:
        if cmd == "scan":
            _print_json(scan_workspace(_confined_dir(path, workspace)))
        elif cmd == "classify":
            text = extract_pdf_text(path, pages=[0], workspace=workspace)
            doc_type = classify_document(text.get(0, ""))
            print(f"Document type: {doc_type}")
        elif cmd == "extract":
            pages = [int(arg3)] if arg3 is not None else None
            text = extract_pdf_text(path, pages=pages, workspace=workspace)
            for pg, content in sorted(text.items()):
                print(f"=== Page {pg} ===")
                print(content)
                print()
        elif cmd == "tables":
            pages = [int(arg3)] if arg3 is not None else None
            _print_json(extract_tables(path, pages=pages, workspace=workspace))
        elif cmd == "csv":
            _print_json(extract_csv(path, workspace=workspace))
        elif cmd == "excel":
            sheets = [arg3] if arg3 is not None else None
            _print_json(extract_excel(path, sheet_names=sheets, workspace=workspace))
        else:
            print(f"Error: Unknown command '{redact(cmd)}'. "
                  f"Valid commands: scan, classify, extract, tables, csv, excel",
                  file=sys.stderr)
            sys.exit(1)
    except (FileNotFoundError, ValueError, NotADirectoryError) as e:
        print(f"Error: {redact(e)}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:  # noqa: BLE001 — the boundary, not a handler
        # Everything else — a page out of range, an off-lock refusal, a parser
        # exception quoting the bytes it choked on — propagated as a full
        # traceback, into the same transcript the masking above exists to keep
        # clean. One redacted line, class named so the failure is still
        # diagnosable, and the traceback stays out of the worker's context.
        print(f"Error: {type(e).__name__}: {redact(e)}", file=sys.stderr)
        sys.exit(1)
