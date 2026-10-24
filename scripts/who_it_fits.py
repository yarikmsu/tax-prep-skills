"""Recompute the figures in docs/WHO-IT-FITS.md from the public files behind them.

Run from the clone's root:

    sh agents/lib/run.sh -m scripts.who_it_fits DATA_DIR
    sh agents/lib/run.sh -m scripts.who_it_fits

With DATA_DIR it reads the downloads below, saves what it measured to
scripts/who_it_fits_figures.json, and prints the page's table and the phrases
its prose has to contain. Without one it prints them from that file. It exits
1 if a comparison the prose makes ("the highest of the five") no longer holds.
tests/test_capabilities.py holds the page to the same output, so a cell or a
figure edited by hand fails.

DATA_DIR holds these downloads, under the names they are published with:

- IRS Statistics of Income, Historic Table 2, tax year 2023, all states:
  `23in55cmcsv.csv`, https://www.irs.gov/statistics/soi-tax-stats-historic-table-2
- IRS Statistics of Income, county data, tax year 2023: `23incyallagi.csv`,
  https://www.irs.gov/statistics/soi-tax-stats-county-data
- IRS Statistics of Income, migration data, 2022 to 2023: `stateinflow2223.csv`,
  `stateoutflow2223.csv`, `countyinflow2223.csv` and `countyoutflow2223.csv`,
  https://www.irs.gov/statistics/soi-tax-stats-migration-data
- Census Bureau, county-to-county commuting flows, ACS 2016-2020, `table1.xlsx`,
  https://www.census.gov/data/tables/2020/demo/metro-micro/commuting-flows-2020.html

The national figures come from Publication 1304, whose tables are `.xls`. The
locked dependencies have no reader for that format, so `PUB1304` holds those
counts by hand, each with its table and column.

A download whose layout has changed stops the run: a missing column raises
KeyError, and a missing summary row or header raises ValueError. Neither is
read as zero.
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

from agents.lib.capabilities import _NO_INCOME_TAX_STATES

TAX_YEAR = 2023
FIGURES = Path(__file__).with_name("who_it_fits_figures.json")
NYC_COUNTIES = frozenset({"005", "047", "061", "081", "085"})

# Persona -> (state FIPS, counties or None for the whole state, whether a move
# between the area and the rest of its state stops the return). New York
# City's does: the City tax refuses a part-year City resident at compute.
PERSONAS = {
    "California": ("06", None, False),
    "Florida": ("12", None, False),
    "New York City": ("36", NYC_COUNTIES, True),
    "North Carolina": ("37", None, False),
    "Seattle (King County)": ("53", frozenset({"033"}), False),
}

# The $500,000+ itemizing share, measured for the five states; the page quotes
# California's.
PERSONA_STATES = ("CA", "FL", "NY", "NC", "WA")

# Publication 1304, tax year 2023, by adjusted gross income.
PUB1304 = {
    # Table 1.4, column 1.
    "returns": {"all": 160_602_107, "200k": 13_534_633, "500k": 2_578_815},
    # Table 1.4, columns 51 and 53: rent net income, and rent net loss
    # including the nondeductible (suspended) loss. Every return with rent.
    "rent_net_income": {"all": 4_886_466, "200k": 1_328_261},
    "rent_net_loss": {"all": 4_336_173, "200k": 1_255_072},
    # Table 1.4, columns 63 and 65: total rental and royalty net income and
    # net loss, the national counterpart of the state field N25870.
    "rental_royalty_net_income": {"all": 6_346_357, "200k": 1_623_243},
    "rental_royalty_net_loss": {"all": 3_196_282, "200k": 429_257},
    # Table 3.3, column 78: household employment tax.
    "household_employment": {"all": 242_743, "200k": 120_158, "500k": 71_061},
    # Table 3.1A, column 1: children's returns with Form 8615, every income.
    "form_8615": {"all": 228_992},
}

BLOCK_ROWS = ("rental_income", "marketplace_insurance", "part_year_residency",
              "move_abroad", "nonresident_source_income")
ROW_LABELS = {
    "rental_income": "Rental real estate, Schedule E (`rental_income`)",
    "marketplace_insurance": "Marketplace health insurance with an advance "
                             "credit, Form 8962 (`marketplace_insurance`)",
    "part_year_residency": "Moved to or from another state during the year "
                           "(`part_year_residency`)",
    "move_abroad": "Moved to or from another country during the year "
                   "(`move_abroad`)",
    "nonresident_source_income": "Works in another state "
                                 "(`nonresident_source_income`) ³",
}


# --- the national figures and the arithmetic ---------------------------------

def pct(share: float) -> str:
    """A percentage to two significant figures, as the page prints it."""
    rounded = float(f"{share:.2g}")
    if rounded >= 10:
        return f"{rounded:.0f}%"
    if rounded >= 1:
        return f"{rounded:.1f}%"
    return f"{rounded:.2f}%"


def as_printed(share: float) -> float:
    return float(pct(share).rstrip("%"))


def national() -> dict[str, float]:
    """The national figures, in percent, before rounding."""
    returns = PUB1304["returns"]

    def share(key: str, group: str) -> float:
        return 100 * PUB1304[key][group] / returns[group]

    out = {}
    for group in ("all", "200k"):
        out[f"any_rent_{group}"] = (share("rent_net_income", group)
                                    + share("rent_net_loss", group))
        out[f"rental_field_{group}"] = (share("rental_royalty_net_income", group)
                                        + share("rental_royalty_net_loss", group))
        # Negative across all returns, where royalty-only returns outnumber
        # the owners the field misses: nothing is added there.
        out[f"rental_gap_{group}"] = max(
            0.0, out[f"any_rent_{group}"] - out[f"rental_field_{group}"])
        out[f"household_{group}"] = share("household_employment", group)
    out["household_500k"] = share("household_employment", "500k")
    out["kiddie_all"] = 100 * PUB1304["form_8615"]["all"] / returns["all"]
    # As if every child belonged to a different $200,000+ household.
    out["kiddie_200k"] = 100 * PUB1304["form_8615"]["all"] / returns["200k"]
    return out


def national_as_printed() -> dict[str, float]:
    raw = national()
    printed = {k: as_printed(v) for k, v in raw.items()}
    for group in ("all", "200k"):
        # The page gives the rental gap in whole points.
        printed[f"rental_gap_{group}"] = round(raw[f"rental_gap_{group}"])
    return printed


def at_least_one(shares: list[float]) -> float:
    """The share hitting at least one block, the blocks taken as independent."""
    missed = 1.0
    for share in shares:
        missed *= 1 - share / 100
    return 100 * (1 - missed)


def at_least_one_ends(column: dict[str, str], printed: dict[str, float]) -> dict[str, str]:
    """One persona's at-least-one range per income group, from its cells as
    printed. The low end is the rows and household employees; the high end adds
    the rental owners the IRS field misses and the Form 8615 figure. One figure
    where both ends round the same."""
    figures: dict[str, list[float]] = {"200k": [], "all": []}
    for rule in BLOCK_ROWS:
        cells = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)%", column[rule])]
        high, every = cells if len(cells) == 2 else cells * 2
        figures["200k"].append(high)
        figures["all"].append(every)
    ends = {}
    for group, shares in figures.items():
        low = at_least_one(shares + [printed[f"household_{group}"]])
        high = at_least_one([shares[0] + printed[f"rental_gap_{group}"]] + shares[1:]
                            + [printed[f"household_{group}"], printed[f"kiddie_{group}"]])
        low_s, high_s = f"{low:.0f}", f"{high:.0f}"
        ends[group] = f"{low_s}%" if low_s == high_s else f"{low_s}–{high_s}%"
    return ends


# --- reading the downloads ----------------------------------------------------

def _num(cell: str) -> float:
    cell = cell.strip().replace('"', "").replace(",", "")
    return float(cell) if cell not in ("", "*", "**") else 0.0


def _dict_rows(path: Path):
    with path.open(encoding="latin-1", newline="") as f:
        yield from csv.DictReader(f)


def _require(present: set, expected: set, where: str) -> None:
    """A bracket or a county missing from a download would still leave a
    denominator, and a plausible share computed from part of the returns."""
    missing = expected - present
    if missing:
        raise ValueError(f"{where}: missing brackets {sorted(missing)}; "
                         "the layout has changed")


def read_county_file(data: Path, personas: dict = PERSONAS) -> tuple[dict, list]:
    """The county file, cut to what is used: each state's FIPS code by postal
    code, and the rows of the persona counties."""
    wanted = {(fips, county) for fips, area, _ in personas.values() if area
              for county in area}
    fips_of: dict[str, str] = {}
    rows = []
    for r in _dict_rows(data / "23incyallagi.csv"):
        fips, county = r["STATEFIPS"].strip(), r["COUNTYFIPS"].strip()
        fips_of[r["STATE"].strip()] = fips
        if (fips, county) in wanted:
            rows.append({"fips": fips, "county": county, "stub": r["agi_stub"].strip(),
                         **{k: _num(r[k]) for k in ("N1", "N25870", "N85775")}})
    return fips_of, rows


def income_shares(data: Path, fips_of: dict, county_rows: list,
                  personas: dict = PERSONAS) -> dict[str, dict[str, float]]:
    """Rental (N25870) and advance premium credit (N85775), in percent of
    returns, at $200,000+ and for every return, per persona. The state file's
    brackets 8-10 are $200,000+ and 0 is every return; the county file's
    bracket 8 is $200,000+, and its eight brackets sum to every return."""
    postal_of = {fips: postal for postal, fips in fips_of.items()}
    wanted = {postal_of[fips] for fips, area, _ in personas.values() if area is None}
    state_rows = [{"state": r["STATE"].strip(), "stub": r["AGI_STUB"].strip(),
                   **{k: _num(r[k]) for k in ("N1", "N25870", "N85775")}}
                  for r in _dict_rows(data / "23in55cmcsv.csv")
                  if r["STATE"].strip() in wanted]
    out = {}
    for persona, (fips, area, _) in personas.items():
        if area is None:
            rows = [r for r in state_rows if r["state"] == postal_of[fips]]
            groups = {"all": {"0"}, "200k": {"8", "9", "10"}}
            _require({r["stub"] for r in rows}, {"0", "8", "9", "10"}, persona)
        else:
            rows = [r for r in county_rows if r["fips"] == fips and r["county"] in area]
            groups = {"all": set("12345678"), "200k": {"8"}}
            _require({(r["county"], r["stub"]) for r in rows},
                     {(county, stub) for county in area for stub in "12345678"}, persona)
        out[persona] = {}
        for group, stubs in groups.items():
            picked = [r for r in rows if r["stub"] in stubs]
            n = sum(r["N1"] for r in picked)
            if not n:
                raise ValueError(f"{persona}: no returns in brackets {sorted(stubs)}")
            out[persona][f"rental_{group}"] = 100 * sum(r["N25870"] for r in picked) / n
            out[persona][f"marketplace_{group}"] = 100 * sum(r["N85775"] for r in picked) / n
    return out


def itemizing_500k(data: Path, states: tuple = PERSONA_STATES) -> dict[str, float]:
    """The share of $500,000+ returns that itemize (N04470), per state."""
    rows = [r for r in _dict_rows(data / "23in55cmcsv.csv")
            if r["AGI_STUB"].strip() in ("9", "10")]
    out = {}
    for state in states:
        picked = [r for r in rows if r["STATE"].strip() == state]
        _require({r["AGI_STUB"].strip() for r in picked}, {"9", "10"}, state)
        out[state] = 100 * sum(_num(r["N04470"]) for r in picked) / sum(
            _num(r["N1"]) for r in picked)
    return out


# Each side of a migration file: which prefix is the home place, which the
# other end of the move.
_SIDES = {"out": ("y1_", "y2_"), "in": ("y2_", "y1_")}
_STATE_LABELS = {"out": ("Non-migrants", "Same State", "US and Foreign", "US", "Foreign"),
                 "in": ("US and Foreign", "US", "Foreign")}
_COUNTY_LABELS = {"out": ("Non-migrants", "US and Foreign", "Different State", "Foreign"),
                  "in": ("US and Foreign", "Different State", "Foreign")}


def _label(name: str) -> str | None:
    """A migration file's summary rows, by name; None for a single place."""
    if "Total Migration-" in name:
        return name.split("Total Migration-")[1]
    return "Non-migrants" if name.endswith("Non-migrants") else None


def read_flows(data: Path, level: str, side: str, home_states: set) -> list[dict]:
    """One migration file's rows whose home state is wanted."""
    home, other = _SIDES[side]
    name = f"{other}state_name" if level == "state" else f"{other}countyname"
    rows = []
    for r in _dict_rows(data / f"{level}{side}flow2223.csv"):
        if r[f"{home}statefips"] not in home_states:
            continue
        rows.append({
            "home_state": r[f"{home}statefips"],
            "home_county": r[f"{home}countyfips"] if level == "county" else None,
            "other_state": r[f"{other}statefips"],
            "other_county": r[f"{other}countyfips"] if level == "county" else None,
            "name": r[name], "n": int(r["n1"])})
    return rows


def _labelled(rows: list[dict], required: tuple, where: str) -> Counter:
    counts: Counter = Counter()
    for r in rows:
        if label := _label(r["name"]):
            counts[label] += r["n"]
    missing = [label for label in required if label not in counts]
    if missing:
        raise ValueError(f"{where}: no {missing} summary rows; the layout has changed")
    return counts


def state_moves(flows: dict, fips: str) -> dict[str, int]:
    """Households leaving or entering a whole state, and the households that
    filed there in either year."""
    c = {side: _labelled([r for r in flows["state", side] if r["home_state"] == fips],
                         _STATE_LABELS[side], f"state{side}flow {fips}")
         for side in ("out", "in")}
    return {
        "households": (c["out"]["Non-migrants"] + c["out"]["Same State"]
                       + c["out"]["US and Foreign"] + c["in"]["US and Foreign"]),
        "state_line": c["out"]["US"] + c["in"]["US"],
        "abroad": c["out"]["Foreign"] + c["in"]["Foreign"],
        "in_state": 0,
    }


def area_moves(flows: dict, fips: str, area: frozenset) -> dict[str, int]:
    """The same for a group of counties. A move between two of them is not a
    move out of the area; one to the rest of the state is, without crossing
    the state line."""
    totals = {side: Counter() for side in ("out", "in")}
    for side in ("out", "in"):
        for county in sorted(area):
            rows = [r for r in flows["county", side]
                    if r["home_state"] == fips and r["home_county"] == county]
            totals[side].update(_labelled(rows, _COUNTY_LABELS[side],
                                          f"county{side}flow {fips}{county}"))
            totals[side]["within area"] += sum(
                r["n"] for r in rows if r["other_state"] == fips
                and r["other_county"] in area and r["other_county"] != county)
    out, inn = totals["out"], totals["in"]
    leaving = out["US and Foreign"] - out["within area"]
    entering = inn["US and Foreign"] - inn["within area"]
    state_line = out["Different State"] + inn["Different State"]
    abroad = out["Foreign"] + inn["Foreign"]
    return {
        "households": out["Non-migrants"] + out["within area"] + leaving + entering,
        "state_line": state_line,
        "abroad": abroad,
        "in_state": leaving + entering - state_line - abroad,
    }


def moves(flows: dict, personas: dict = PERSONAS) -> dict[str, dict[str, float]]:
    """Per persona, in percent of households: moves across the state line
    (abroad left out), moves abroad, and moves to or from the rest of the
    state. The migration files cover every income together."""
    out = {}
    for persona, (fips, area, _) in personas.items():
        counts = state_moves(flows, fips) if area is None else area_moves(flows, fips, area)
        out[persona] = {k: 100 * counts[k] / counts["households"]
                        for k in ("state_line", "abroad", "in_state")}
    return out


def arrivals(flows: dict, personas: dict = PERSONAS) -> dict[str, list]:
    """Per whole-state persona, the three states most arrivals came from, as
    [name, percent of arrivals from another state]."""
    out = {}
    for persona, (fips, area, _) in personas.items():
        if area is not None:
            continue
        origins: Counter = Counter()
        for r in flows["state", "in"]:
            if r["home_state"] == fips and r["other_state"] != fips \
                    and r["other_state"].isdigit() and int(r["other_state"]) <= 56:
                origins[r["name"]] += r["n"]
        total = sum(origins.values())
        out[persona] = [[name, 100 * n / total] for name, n in origins.most_common(3)]
    return out


COMMUTING_HEADER = ("State FIPS Code", "County FIPS Code", "State Name", "County Name",
                    "State FIPS Code", "County FIPS Code", "State Name", "County Name",
                    "Workers in Commuting Flow")


def _guard_xml() -> None:
    """AGENTS.md section 7: a downloaded .xlsx is parsed through defusedxml, or
    not at all."""
    try:
        import defusedxml
    except ImportError:
        raise SystemExit("who_it_fits: defusedxml is not installed, so the .xlsx "
                         "cannot be read safely. Run through agents/lib/run.sh, "
                         "whose venv has the locked requirements.")
    defusedxml.defuse_stdlib()


def commuting(data: Path, no_tax_fips: set, personas: dict = PERSONAS) -> dict:
    """Per persona, the percent of WORKERS whose workplace is in another state
    whose FIPS code is not in `no_tax_fips`; and the three commonest
    out-of-state workplace states, as [name, percent of out-of-state workers,
    whether it blocks]. A workplace outside the fifty states and DC, abroad or
    in Puerto Rico, does not block and stays in the denominator."""
    _guard_xml()
    from openpyxl import load_workbook

    rows = load_workbook(data / "table1.xlsx", read_only=True).active.iter_rows(
        values_only=True)
    for row in rows:
        if tuple(row[:9]) == COMMUTING_HEADER:
            break
    else:
        raise ValueError("table1.xlsx: no header row; the layout has changed")
    workers: Counter = Counter()
    blocking: Counter = Counter()
    places: dict[str, Counter] = {p: Counter() for p in personas}
    for row in rows:
        home_state, home_county, _, _, work_state, _, work_name, _, n = row[:9]
        if not str(home_state).isdigit() or not isinstance(n, (int, float)):
            continue
        state = str(work_state).strip()
        abroad = not state.isdigit() or int(state) > 56
        for persona, (fips, area, _) in personas.items():
            if int(home_state) != int(fips) or (area is not None and home_county not in area):
                continue
            workers[persona] += n
            if abroad or int(state) == int(fips):
                continue
            blocks = f"{int(state):02d}" not in no_tax_fips
            places[persona][work_name, blocks] += n
            if blocks:
                blocking[persona] += n
    return {
        "blocking": {p: 100 * blocking[p] / workers[p] for p in personas},
        "workplaces": {p: [[name, 100 * n / sum(where.values()), blocks]
                           for (name, blocks), n in where.most_common(3)]
                       for p, where in places.items()},
    }


def no_tax_fips(fips_of: dict) -> set:
    """The gate's no-income-tax states by FIPS code, from the IRS's own
    postal-to-FIPS pairs rather than a list typed here."""
    return {fips_of[postal] for postal in _NO_INCOME_TAX_STATES}


def ca_scope_shares(data: Path, postal: str = "CA") -> dict:
    """CA-specific stops/bounds from Historic Table 2, not independent totals.

    N02500 is federally taxable Social Security; N00700 taxable state/local
    refunds; N02300 unemployment; N00400 ALL tax-exempt interest, which only
    bounds the non-CA municipal subset from above.
    """
    fields = {"taxable_ss": "N02500", "taxable_refund": "N00700",
              "unemployment": "N02300", "muni_upper_bound": "N00400"}
    rows = [r for r in _dict_rows(data / "23in55cmcsv.csv") if r["STATE"].strip() == postal]
    by_stub = {r["AGI_STUB"].strip(): r for r in rows}
    if len(by_stub) != len(rows):
        raise ValueError("CA scope statistics: duplicate AGI brackets")
    _require(set(by_stub), {"0", "8", "9", "10"}, "CA scope statistics")
    out = {}
    for group, stubs in (("all", ("0",)), ("200k", ("8", "9", "10"))):
        denominator = sum(_num(by_stub[s]["N1"]) for s in stubs)
        if not math.isfinite(denominator) or denominator <= 0:
            raise ValueError("CA scope statistics: denominator must be positive")
        for name, field in fields.items():
            numerator = sum(_num(by_stub[s][field]) for s in stubs)
            if not math.isfinite(numerator) or not 0 <= numerator <= denominator:
                raise ValueError("CA scope statistics: count outside return population")
            out[f"{name}_{group}"] = 100 * numerator / denominator
    return out


def extract(data: Path) -> dict:
    """Everything the page quotes that comes from the downloads."""
    fips_of, county_rows = read_county_file(data)
    home_states = {fips for fips, _, _ in PERSONAS.values()}
    flows = {(level, side): read_flows(data, level, side, home_states)
             for level in ("state", "county") for side in ("out", "in")}
    work = commuting(data, no_tax_fips(fips_of))
    return {
        "tax_year": TAX_YEAR,
        "income": income_shares(data, fips_of, county_rows),
        "ca_scope": ca_scope_shares(data),
        "itemizing_500k": itemizing_500k(data),
        "moves": moves(flows),
        "arrivals": arrivals(flows),
        "commuting": work["blocking"],
        "workplaces": work["workplaces"],
    }


# --- rendering the page -------------------------------------------------------

def _moving(figures: dict, persona: str) -> float:
    """The move figure the table prints: across the state line, plus moves to
    or from the rest of the state where those stop the return too."""
    m = figures["moves"][persona]
    return m["state_line"] + (m["in_state"] if PERSONAS[persona][2] else 0)


def columns(figures: dict) -> dict[str, dict[str, str]]:
    """Each persona's cells in the page's table, above the at-least-one row."""
    out = {}
    for persona, (_, _, local_moves_block) in PERSONAS.items():
        s = figures["income"][persona]
        out[persona] = {
            "rental_income": f"{pct(s['rental_200k'])} · {pct(s['rental_all'])}",
            "marketplace_insurance":
                f"{pct(s['marketplace_200k'])} · {pct(s['marketplace_all'])}",
            "part_year_residency": pct(_moving(figures, persona))
                                   + ("&nbsp;¹" if local_moves_block else ""),
            "move_abroad": pct(figures["moves"][persona]["abroad"]),
            "nonresident_source_income": pct(figures["commuting"][persona]),
        }
    return out


def ranges(figures: dict) -> dict[str, dict[str, str]]:
    """Each persona's at-least-one range, per income group."""
    printed = national_as_printed()
    return {p: at_least_one_ends(column, printed) for p, column in columns(figures).items()}


def render_table(figures: dict) -> list[str]:
    """The page's table, row for row."""
    cells, ends = columns(figures), ranges(figures)
    lines = ["| Block (gate rule) | " + " | ".join(PERSONAS) + " |",
             "|---" * (len(PERSONAS) + 1) + "|"]
    for rule in BLOCK_ROWS:
        lines.append(f"| {ROW_LABELS[rule]} | "
                     + " | ".join(cells[p][rule] for p in PERSONAS) + " |")
    lines.append("| **At least one block** ² | " + " | ".join(
        f"**{ends[p]['200k']}** · {ends[p]['all']}" for p in PERSONAS) + " |")
    return lines


def _span(ranges_: list[str]) -> str:
    numbers = [float(x) for r in ranges_ for x in re.findall(r"\d+", r)]
    return f"{min(numbers):.0f}–{max(numbers):.0f}%"


def _low(range_: str) -> float:
    return float(re.findall(r"\d+", range_)[0])


def phrases(figures: dict) -> list[str]:
    """What the page's prose says, word for word, given these figures. The test
    compares with whitespace collapsed, so a phrase may wrap across lines."""
    ca, fl, ny, nc, sea = PERSONAS
    n, gap = {k: pct(v) for k, v in national().items()}, national_as_printed()
    inc = {p: {k: pct(v) for k, v in s.items()} for p, s in figures["income"].items()}
    com = {p: pct(v) for p, v in figures["commuting"].items()}
    moved = {p: pct(_moving(figures, p)) for p in PERSONAS}
    ends, work = ranges(figures), figures["workplaces"]
    itemize = figures["itemizing_500k"]
    origin = figures["arrivals"][fl][0]
    return [
        "common federal/residency blocks alone reach "
        f"{_span([e['all'] for e in ends.values()])} of returns, depending "
        f"on the persona, and {_span([e['200k'] for e in ends.values()])} above $200,000",
        f"Of the {moved[ny]}, {pct(figures['moves'][ny]['state_line'])} cross a state line",
        f"{pct(figures['moves'][ny]['in_state'])} move between the City and the rest "
        "of New York State",
        f"({gap['rental_gap_200k']} points nationally above $200,000, none across all "
        "returns)",
        f"{n['household_200k']} of $200,000+ returns, {n['household_500k']} of "
        f"$500,000+, and {n['household_all']} of all returns",
        f"about {round(PUB1304['form_8615']['all'], -3):,} children's returns used "
        f"Form 8615 for {TAX_YEAR}, {n['kiddie_all']} of all returns",
        f"that would be {n['kiddie_200k']} of them",
        f"Above $200,000, {inc[ca]['rental_200k']} of returns carry rental or royalty "
        f"income, and that one block is most of the {ends[ca]['200k']} that hit any",
        f"Moves ({moved[ca]}) are the lowest of the five",
        f"{work[ca][0][0]}, the commonest out-of-state workplace",
        f"{pct(itemize['CA'])} of California returns above $500,000 itemize, and for "
        "them the federal return applies the SALT cap's phasedown",
        f"{inc[fl]['marketplace_all']} of Florida returns took an advance premium credit",
        f"highest of the five for all returns ({ends[fl]['all']})",
        f"Above $200,000 that falls to {inc[fl]['marketplace_200k']}, and rental takes over",
        f"{moved[fl]} of households moved to or from another state during the year",
        f"the largest group of arrivals, {pct(origin[1])}, comes from {origin[0]}",
        f"Above $200,000, {inc[ny]['rental_200k']} carry rental income, the lowest "
        "rental rate of the five",
        f"Moves are the highest: {moved[ny]} of City households moved in or out in a year",
        f"{com[ny]} of City workers work in another state that blocks",
        f"{pct(work[ny][0][1])} of the City's out-of-state commuters work in {work[ny][0][0]}",
        f"Together, {ends[ny]['200k']} of $200,000+ returns hit a block",
        f"come out as Florida's do, {ends[nc]['200k']} hitting at least one",
        f"fewer are blocked, {ends[nc]['all']} against Florida's {ends[fl]['all']}",
        f"fewer took an advance premium credit: {inc[nc]['marketplace_all']} against "
        f"{inc[fl]['marketplace_all']}",
        f"{moved[nc]} of households moved to or from another state, and {com[nc]} of "
        "workers work in another state that blocks",
        f"{pct(work[nc][0][1] + work[nc][1][1])} of North Carolina's out-of-state "
        f"commuters work in {work[nc][0][0]} or {work[nc][1][0]}",
        f"Rental is {inc[nc]['rental_200k']} above $200,000",
        f"rental at {inc[sea]['rental_200k']} above $200,000, moves across the state "
        f"line at {moved[sea]}, and almost no commuting ({com[sea]})",
        f"{ends[sea]['200k']} of $200,000+ returns hit one",
        f"returns reporting any rent are {n['any_rent_200k']} of $200,000+ returns "
        f"against {n['rental_field_200k']} by this field",
    ] + ca_scope_phrases(figures)


def ca_scope_phrases(figures: dict) -> list[str]:
    s = figures["ca_scope"]
    return [
        f"Taxable Social Security: {s['taxable_ss_200k']:.1f}% above $200,000 and {s['taxable_ss_all']:.1f}% of all California returns",
        f"Taxable state/local refunds: {s['taxable_refund_200k']:.1f}% above $200,000 and {s['taxable_refund_all']:.1f}% overall",
        f"Unemployment compensation: {s['unemployment_200k']:.1f}% above $200,000 and {s['unemployment_all']:.1f}% overall",
        f"All tax-exempt interest bounds the non-CA municipal subset from above: {s['muni_upper_bound_200k']:.1f}% above $200,000 and {s['muni_upper_bound_all']:.1f}% overall",
    ]


def claims(figures: dict) -> dict[str, bool]:
    """The comparisons the page's prose makes, which a newer year can undo."""
    ca, fl, ny, nc, _ = PERSONAS
    inc, com, work = figures["income"], figures["commuting"], figures["workplaces"]
    moved = {p: _moving(figures, p) for p in PERSONAS}
    ends = ranges(figures)
    low_all = {p: _low(ends[p]["all"]) for p in PERSONAS}
    others = [p for p in PERSONAS]
    return {
        "California has the lowest common-block all-returns share":
            min(others, key=low_all.get) == ca,
        "rental is the largest of the common measured blocks above $200,000 in every persona": all(
            inc[p]["rental_200k"] > max(inc[p]["marketplace_200k"], moved[p], com[p])
            for p in PERSONAS),
        "California has the fewest moves": min(others, key=moved.get) == ca,
        "California's commonest out-of-state workplace does not block": not work[ca][0][2],
        "Florida has the highest all-returns marketplace share":
            max(others, key=lambda p: inc[p]["marketplace_all"]) == fl,
        "the marketplace credit is Florida's largest all-returns block":
            inc[fl]["marketplace_all"] > max(inc[fl]["rental_all"], moved[fl], com[fl]),
        "Florida has the highest all-returns share": max(others, key=low_all.get) == fl,
        "New York City has the lowest rental share above $200,000":
            min(others, key=lambda p: inc[p]["rental_200k"]) == ny,
        "New York City has the most moves": max(others, key=moved.get) == ny,
        "North Carolina's $200,000+ range is Florida's": ends[nc]["200k"] == ends[fl]["200k"],
        "North Carolina's two commonest out-of-state workplaces both block":
            work[nc][0][2] and work[nc][1][2],
        "the rental field misses no one across all returns":
            national()["rental_gap_all"] == 0,
    }


def main(argv: list[str]) -> int:
    if len(argv) > 2:
        print(__doc__, file=sys.stderr)
        return 2
    if len(argv) == 2:
        figures = extract(Path(argv[1]))
        FIGURES.write_text(json.dumps(figures, indent=1, sort_keys=True) + "\n",
                           encoding="utf-8")
    figures = json.loads(FIGURES.read_text(encoding="utf-8"))
    print("\n".join(render_table(figures) + [""] + phrases(figures)))
    failed = [claim for claim, holds in claims(figures).items() if not holds]
    for claim in failed:
        print(f"no longer true: {claim}", file=sys.stderr)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
